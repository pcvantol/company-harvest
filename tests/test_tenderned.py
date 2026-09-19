"""Offline TenderNed-bronproeven met uitsluitend synthetische bedrijven."""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from test_pre_kvk import _full_sources

from company_lookup import tenderned
from company_lookup.core import HTTP_USER_AGENT, HarvestError, Run, read_tsv
from company_lookup.pre_kvk import CURRENT_SOURCE_IDS, build_pre_kvk_list, source_scope
from company_lookup.pre_kvk_filter import build_pre_kvk_filter
from company_lookup.sources import read_catalog

XLSX_URL = (
    "https://www.tenderned.nl/cms/sites/default/files/2026-07/"
    "Dataset_Tenderned-compleet-2021-01-01-2026-06-30.xlsx"
)
JSON_URL = (
    "https://www.tenderned.nl/cms/sites/default/files/2026-07/"
    "Dataset_Tenderned-2026-01-01-2026-06-30.json"
)


def snapshots(root: Path) -> tuple[Path, Path]:
    xlsx = root / "publicaties.xlsx"
    json_file = root / "gunningen.json"
    book = Workbook()
    sheet = book.active
    sheet.title = "OpenData sheet"
    sheet.append([
        "ID publicatie", "Publicatiedatum", "Naam gegunde onderneming",
        "ON kvknummer", "ON land", "ON plaats", "URL TenderNed",
    ])
    for year in range(2021, 2025):
        sheet.append([year, f"{year}-01-01", None, None, None, None, None])
    sheet.append([1, "2025-03-01", "Oude Leverancier B.V.", "12345678", "Nederland", "Utrecht", ""])
    sheet.append([2, "2025-06-01", "Buitenland GmbH", "87654321", "Duitsland", "Bonn", ""])
    sheet.append([3, "2026-02-01", "Eerste Gunning B.V.", "23456789", "Nederland", "Arnhem", ""])
    book.save(xlsx)
    json_file.write_text(json.dumps({"releases": [{
        "id": "3", "date": "2026-02-01", "awards": [{"suppliers": [
            {"id": "23456789", "name": "Eerste Gunning B.V."},
            {"id": "56789012", "name": "Tweede Gunning B.V."},
            {"id": "", "name": "Zonder Nummer B.V."},
            {"id": "45678901", "name": "Foreign Ltd"},
        ]}],
        "parties": [
            {"id": "23456789", "name": "Eerste Gunning B.V.", "roles": ["supplier"],
             "address": {"countryName": "Nederland", "locality": "Arnhem"}},
            {"id": "56789012", "name": "Tweede Gunning B.V.", "roles": ["supplier"],
             "address": {"countryName": "Nederland", "locality": "Utrecht"}},
            {"id": "", "name": "Zonder Nummer B.V.", "roles": ["supplier"],
             "address": {"countryName": "Nederland", "locality": "Den Haag"}},
            {"id": "45678901", "name": "Foreign Ltd", "roles": ["supplier"],
             "address": {"countryName": "Duitsland", "locality": "Bonn"}},
        ],
    }]}), encoding="utf-8")
    return xlsx, json_file


def mock_official_download(
    monkeypatch: pytest.MonkeyPatch, xlsx: Path, json_file: Path,
) -> None:
    def fake_download(url: str, destination: Path, _maximum: int) -> dict[str, object]:
        if url == tenderned.PAGE_URL:
            destination.write_text(
                f"<a href='{XLSX_URL}'>XLSX</a><a href='{JSON_URL}'>JSON</a>",
                encoding="utf-8",
            )
        elif url == XLSX_URL:
            shutil.copyfile(xlsx, destination)
        elif url == JSON_URL:
            shutil.copyfile(json_file, destination)
        else:
            raise AssertionError(f"unexpected TenderNed URL: {url}")
        return {"url": url, "http_status": 200, "downloaded_bytes": destination.stat().st_size}

    monkeypatch.setattr(tenderned, "_download", fake_download)


def test_select_downloads_requires_same_snapshot_and_official_host() -> None:
    page = """<a href='/cms/sites/default/files/2026-07/Dataset_Tenderned-compleet-2021-01-01-2026-06-30.xlsx'>Excel</a>
    <a href='/cms/sites/default/files/2026-07/Dataset_Tenderned-2026-01-01-2026-06-30.json'>JSON</a>
    <a href='https://evil.example/Dataset_Tenderned-2026-01-01-2026-06-30.json'>evil</a>"""
    xlsx, json_url, year = tenderned._select_downloads(page)
    assert xlsx.startswith("https://www.tenderned.nl/")
    assert json_url.startswith("https://www.tenderned.nl/") and year == 2026
    with pytest.raises(HarvestError, match="JSON"):
        tenderned._select_downloads(page.replace("2026-06-30.json", "2026-03-31.json"))
    with pytest.raises(HarvestError, match="XLSX"):
        tenderned._select_downloads("<a href='https://evil.example/file.xlsx'>bad</a>")


def test_local_adapter_preserves_kvk_hints_and_rejects_foreign(
    run: Run, tmp_path: Path,
) -> None:
    xlsx, json_file = snapshots(tmp_path)
    source, rejected, report_path, _ = tenderned.collect_tenderned(run, xlsx, json_file)
    candidates = read_tsv(source)
    report = json.loads(report_path.read_text())
    assert {row["original_name"] for row in candidates} == {
        "Oude Leverancier B.V.", "Eerste Gunning B.V.", "Tweede Gunning B.V.",
        "Zonder Nummer B.V.",
    }
    assert {row["source_kvk_hint"] for row in candidates} == {
        "12345678", "23456789", "56789012", "",
    }
    assert report["counts"]["supplier_observations"] == 6
    assert report["counts"]["candidate_records"] == 4
    assert report["counts"]["excluded_records"] == 2
    assert report["counts"]["xlsx_latest_year_superseded"] == 1
    assert report["counts"]["xlsx_without_awarded_name"] == 4
    assert report["count_closure"] == {"supplier_partition": "CLOSED"}
    assert report["provenance_status"] == "USER_SUPPLIED_OFFICIAL_FORMAT_UNVERIFIED"
    assert next(row for row in read_catalog(run) if row["source_id"] == "tenderned_awards")[
        "live_measurement_status"
    ] == "MEASURED_LOCAL_EVIDENCE"
    assert any(row["reason"] == "MISSING_KVK" for row in read_tsv(rejected))
    assert any(row["reason"] == "NON_NL_OR_UNKNOWN_COUNTRY" for row in read_tsv(rejected))
    assert tenderned.collect_tenderned(run, xlsx, json_file)[0] == source
    assert run.latest_artifact("02", "source_tenderned_awards") == source


def test_five_source_master_and_filter_keep_tenderned_trace(
    run: Run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _full_sources(run)
    xlsx, json_file = snapshots(tmp_path)
    mock_official_download(monkeypatch, xlsx, json_file)
    tenderned.collect_tenderned(run)
    run.record_config("pre_kvk_sources", list(CURRENT_SOURCE_IDS))
    assert source_scope(run) == CURRENT_SOURCE_IDS
    master, report_path = build_pre_kvk_list(run)
    report = json.loads(report_path.read_text())
    assert report["scope"] == list(CURRENT_SOURCE_IDS)
    assert report["sources"]["tenderned_awards"]["candidate_records"] == 4
    assert any("tenderned_awards" in row["source_relations"] for row in read_tsv(master))
    eligible, excluded, _ = build_pre_kvk_filter(run)
    assert any(row["original_name"] == "Tweede Gunning B.V." for row in read_tsv(eligible))
    assert any(row["original_name"] == "Zonder Nummer B.V." for row in read_tsv(excluded))


def test_download_reuse_checks_page_evidence(
    run: Run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    xlsx, json_file = snapshots(tmp_path)
    mock_official_download(monkeypatch, xlsx, json_file)
    original_download = tenderned._download
    downloads: list[str] = []

    def counted_download(url: str, destination: Path, maximum: int) -> dict[str, object]:
        downloads.append(url)
        return original_download(url, destination, maximum)

    monkeypatch.setattr(tenderned, "_download", counted_download)
    first = tenderned.collect_tenderned(run)
    assert len(downloads) == 3
    assert tenderned.collect_tenderned(run) == first
    assert len(downloads) == 3
    page = run.latest_artifact("02", "evidence_tenderned_awards_page")
    assert page is not None
    page.write_text("changed", encoding="utf-8")
    tenderned.collect_tenderned(run)
    assert len(downloads) == 6
    restored = run.latest_artifact("02", "evidence_tenderned_awards_page")
    assert restored is not None and "Dataset_Tenderned" in restored.read_text()


def test_adapter_rejects_missing_pair_bad_country_and_broken_evidence(
    run: Run, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    xlsx, json_file = snapshots(tmp_path)
    with pytest.raises(HarvestError, match="beide"):
        tenderned.collect_tenderned(run, xlsx=xlsx)
    with pytest.raises(HarvestError, match="ontbreekt"):
        tenderned.collect_tenderned(run, xlsx=tmp_path / "missing.xlsx", json_file=json_file)
    source, _, report_path, _ = tenderned.collect_tenderned(run, xlsx, json_file)
    assert json.loads(report_path.read_text())["scope"] == "LOCAL_UNVERIFIED"
    _full_sources(run)
    run.record_config("pre_kvk_sources", list(CURRENT_SOURCE_IDS))
    assert source.is_file()
    with pytest.raises(HarvestError, match="full-archive"):
        build_pre_kvk_list(run)
    mock_official_download(monkeypatch, xlsx, json_file)
    tenderned.collect_tenderned(run, refresh=True)
    evidence = run.latest_artifact("02", "evidence_tenderned_awards_json")
    assert evidence is not None
    evidence.write_text("changed", encoding="utf-8")
    with pytest.raises(HarvestError, match="afwijkende evidence"):
        build_pre_kvk_list(run)


def test_json_requires_actual_award_relation(run: Run, tmp_path: Path) -> None:
    xlsx, json_file = snapshots(tmp_path)
    data = json.loads(json_file.read_text())
    data["releases"][0]["awards"] = []
    json_file.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(HarvestError, match="gunningsrelatie"):
        tenderned.collect_tenderned(run, xlsx, json_file)
    assert run.latest_artifact("02", "source_tenderned_awards") is None


def test_json_rejects_inconsistent_year(run: Run, tmp_path: Path) -> None:
    xlsx, json_file = snapshots(tmp_path)
    data = json.loads(json_file.read_text())
    older_release = dict(data["releases"][0])
    older_release["date"] = "2025-01-01"
    data["releases"].append(older_release)
    json_file.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(HarvestError, match="laatste XLSX-jaar"):
        tenderned.collect_tenderned(run, xlsx, json_file)


def test_json_accepts_publication_without_award(run: Run, tmp_path: Path) -> None:
    xlsx, json_file = snapshots(tmp_path)
    data = json.loads(json_file.read_text())
    data["releases"].append({
        "id": "4", "date": "2026-02-02", "awards": None, "parties": [],
    })
    json_file.write_text(json.dumps(data), encoding="utf-8")
    _, _, report_path, _ = tenderned.collect_tenderned(run, xlsx, json_file)
    assert json.loads(report_path.read_text())["counts"]["json_release_rows"] == 2


def test_download_is_bounded_redirected_and_identified(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    original = httpx.Client
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        assert request.headers["user-agent"] == HTTP_USER_AGENT
        if request.url.path == "/start":
            return httpx.Response(302, headers={"location": "/final"})
        return httpx.Response(200, content=b"data", headers={"content-length": "4"})

    def fake_client(**kwargs: object) -> httpx.Client:
        return original(transport=httpx.MockTransport(handler), **kwargs)

    monkeypatch.setattr(tenderned.httpx, "Client", fake_client)
    target = tmp_path / "download.dat"
    result = tenderned._download("https://www.tenderned.nl/start", target, 10)
    assert target.read_bytes() == b"data"
    assert result["redirects"] == 1 and result["downloaded_bytes"] == 4
    assert len(seen) == 2


@pytest.mark.parametrize(
    ("response", "message"),
    [
        (httpx.Response(302, headers={"location": "https://evil.example/file"}), "officiële"),
        (httpx.Response(200, content=b"123456", headers={"content-length": "6"}), "groottegate"),
        (httpx.Response(200, content=b"x", headers={"content-length": "abc"}), "Content-Length"),
        (httpx.Response(403, content=b"no"), "HTTP 403"),
        (httpx.Response(200, content=b""), "leeg"),
    ],
)
def test_download_rejects_unsafe_or_failed_response(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    response: httpx.Response, message: str,
) -> None:
    original = httpx.Client
    monkeypatch.setattr(
        tenderned.httpx, "Client",
        lambda **kwargs: original(transport=httpx.MockTransport(lambda _request: response), **kwargs),
    )
    with pytest.raises(HarvestError, match=message):
        tenderned._download("https://www.tenderned.nl/start", tmp_path / "result", 4)
    assert not (tmp_path / "result").exists()


def test_download_rejects_body_past_declared_size(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    original = httpx.Client
    monkeypatch.setattr(
        tenderned.httpx, "Client",
        lambda **kwargs: original(
            transport=httpx.MockTransport(
                lambda _request: httpx.Response(200, content=b"12345")
            ), **kwargs,
        ),
    )
    with pytest.raises(HarvestError, match="groottegate"):
        tenderned._download("https://www.tenderned.nl/file", tmp_path / "result", 4)


def test_archive_and_schema_gates_are_fail_closed(run: Run, tmp_path: Path) -> None:
    xlsx, json_file = snapshots(tmp_path)
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(HarvestError, match="ZIP"):
        tenderned._safe_xlsx(broken)
    unsafe = tmp_path / "unsafe.xlsx"
    with zipfile.ZipFile(unsafe, "w") as bundle:
        bundle.writestr("../outside.xml", "x")
    with pytest.raises(HarvestError, match="onveilige"):
        tenderned._safe_xlsx(unsafe)
    unsafe_windows = tmp_path / "unsafe-windows.xlsx"
    with zipfile.ZipFile(unsafe_windows, "w") as bundle:
        bundle.writestr("..\\outside.xml", "x")
    with pytest.raises(HarvestError, match="onveilige"):
        tenderned._safe_xlsx(unsafe_windows)
    data = json.loads(json_file.read_text())
    data["releases"] = []
    json_file.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(HarvestError, match="mist publicaties"):
        tenderned.collect_tenderned(run, xlsx, json_file)


def test_archive_requires_every_year(run: Run, tmp_path: Path) -> None:
    xlsx, json_file = snapshots(tmp_path)
    book = load_workbook(xlsx)
    sheet = book["OpenData sheet"]
    sheet.delete_rows(2, 4)
    book.save(xlsx)
    with pytest.raises(HarvestError, match="mist een jaar"):
        tenderned.collect_tenderned(run, xlsx, json_file)
