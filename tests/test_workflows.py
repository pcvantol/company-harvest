import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from company_harvest.audit import _overlap, trace, verify
from company_harvest.core import HarvestError, initialize_run, read_tsv, sha256, write_tsv
from company_harvest.workflow import (
    _peak_memory,
    active_only,
    consolidate,
    exclude_sole_proprietorships,
    export,
    merge_candidates,
    outcome_metrics,
    report,
    write_xlsx,
)


def _register(run, step, kind, headers, rows):
    path = run.artifact_path(step, kind, "csv")
    write_tsv(path, headers, rows)
    run.register_artifact(path, step, kind)
    return path


def test_harvest_offline_pipeline(run) -> None:
    headers = ["original_name", "country", "nl_evidence", "employees_raw", "employees_date", "employees_scope", "website", "sector", "source_kvk_hint", "source_id", "source_url", "fetched_at", "source_row", "evidence_reference"]
    base = {"country": "Nederland", "nl_evidence": "register", "employees_raw": "", "employees_date": "", "employees_scope": "", "website": "https://alpha.example", "sector": "Bouw", "source_url": "https://example.invalid", "fetched_at": "now", "evidence_reference": "local"}
    rows = [{**base, "original_name": "Alpha B.V.", "source_kvk_hint": "01234567", "source_id": "a", "source_row": "1"}, {**base, "original_name": " Alpha   B.V. ", "source_kvk_hint": "01234567", "source_id": "a", "source_row": "2"}]
    _register(run, "02", "source_a", headers, rows)
    candidates, decisions, conflicts = merge_candidates(run)
    assert len(read_tsv(candidates)) == 1 and read_tsv(decisions)[0]["decision"] == "MERGED_IDENTICAL"
    assert read_tsv(conflicts) == []
    candidate_id = read_tsv(candidates)[0]["candidate_id"]
    match_headers = ["candidate_id", "Bedrijfsnaam", "KVK-nummer", "raw_legal_form", "raw_status", "city", "country", "match_method", "provider", "checked_at", "response_json", "source_relations"]
    public_hit = {"naam": "Alpha B.V.", "kvkNummer": "01234567", "rechtsvormCode": "BV",
                  "actief": True, "inschrijvingsdatum": "20200101", "activiteitomschrijving": "Bouwen",
                  "bezoeklocatie": {"straat": "Dorpsstraat", "huisnummer": 1, "postcode": "1234AB", "plaats": "Utrecht"},
                  "huidigeHandelsNamen": ["Alpha", "Alpha Bouw"], "id": "technical", "bron": "technical", "set": "technical"}
    _register(run, "04", "kvk_matches", match_headers, [{"candidate_id": candidate_id, "Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "01234567", "raw_legal_form": "Besloten vennootschap", "raw_status": "Actief", "city": "Utrecht", "country": "Nederland", "match_method": "SOURCE_KVK_CONFIRMED", "provider": "mock", "checked_at": "now", "response_json": json.dumps([public_hit]), "source_relations": "[]"}])
    _register(run, "05", "kvk_unresolved", ["candidate_id", "reason"], [])
    canonical = consolidate(run)
    assert read_tsv(canonical)[0]["KVK-nummer"] == "01234567"
    assert len(read_tsv(exclude_sole_proprietorships(run)[0])) == 1
    assert len(read_tsv(active_only(run)[0])) == 1
    outputs = export(run)
    assert load_workbook(outputs[1])["Bedrijven"]["B2"].value == "01234567"
    light = read_tsv(outputs[2])
    assert len(light) == 1
    assert {key: light[0][key] for key in (
        "Bedrijfsnaam", "KVK-nummer", "Rechtsvorm (KVK)", "Status (KVK)",
        "Plaats (KVK)", "Land (KVK)", "Website (bron)", "Sector (bron)",
    )} == {
        "Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "01234567",
        "Rechtsvorm (KVK)": "Besloten vennootschap", "Status (KVK)": "Actief",
        "Plaats (KVK)": "Utrecht", "Land (KVK)": "Nederland",
        "Website (bron)": "https://alpha.example",
        "Sector (bron)": "Bouw",
    }
    assert light[0]["Rechtsvormcode (KVK)"] == "BV"
    assert light[0]["Actief-vlag (KVK)"] == "Ja"
    assert light[0]["Inschrijfdatum (KVK)"] == "20200101"
    assert light[0]["Activiteitomschrijving (KVK)"] == "Bouwen"
    assert light[0]["Straat bezoekadres (KVK)"] == "Dorpsstraat"
    assert light[0]["Handelsnamen (KVK)"] == "Alpha; Alpha Bouw"
    assert not any(key in light[0] for key in ("response_json", "source_relations", "id", "bron", "set", "provider", "checked_at"))
    light_sheet = load_workbook(outputs[3])["Bedrijven"]
    assert light_sheet["B2"].value == "01234567"
    assert "response_json" not in [cell.value for cell in light_sheet[1]]
    assert "source_relations" not in [cell.value for cell in light_sheet[1]]
    assert json.loads(outputs[-1].read_text())["schema"] == 2
    assert report(run).is_file()
    outcome = outcome_metrics(run)
    assert outcome["http_user_agent"] == "company-lookup/0.1"
    assert outcome["count_closure"]["status"] == "COMPLETE_CLOSED"
    assert run.latest_artifact("08", "outcome_report").is_file()
    assert verify(run)["valid"]
    old_manifest = outputs[-1]
    old_payload = json.loads(old_manifest.read_text())
    del old_payload["selection_policy"]
    del old_payload["active_rows"]
    old_manifest.write_text(json.dumps(old_payload))
    with run.connect() as connection:
        connection.execute(
            "UPDATE artifacts SET sha256=?,size=? WHERE path=?",
            (sha256(old_manifest), old_manifest.stat().st_size, str(old_manifest.relative_to(run.path))),
        )
    run.record_config("export", {"limit": 100, "allow_partial": False})
    assert verify(run)["valid"]  # oude v2-outputsets zonder selectiepolicy blijven auditbaar
    with run.connect() as connection:
        connection.execute(
            "INSERT INTO kvk_requests(candidate_id,query,state,attempt,request_fingerprint,provider) VALUES(?,?,'SUCCEEDED',1,?,?)",
            (candidate_id, "Alpha B.V.", "fingerprint", "mock"),
        )
    traced = trace(run, "01234567")
    assert traced["trace"] and traced["request_journal"]
    bad_manifest = run.artifact_path("08", "bad_outputset_manifest", "json")
    bad_manifest.write_text('{"schema":1,"files":[]}')
    run.register_artifact(bad_manifest, "08", "outputset_manifest")
    with pytest.raises(HarvestError):
        verify(run)


def test_filters_partial_and_integrity(run) -> None:
    headers = ["Bedrijfsnaam", "KVK-nummer", "raw_legal_form", "raw_status"]
    _register(run, "05", "canonical", headers, [{"Bedrijfsnaam": "Een", "KVK-nummer": "11111111", "raw_legal_form": "eenmanszaak", "raw_status": "actief"}, {"Bedrijfsnaam": "Onbekend", "KVK-nummer": "22222222", "raw_legal_form": "mystery", "raw_status": ""}, {"Bedrijfsnaam": "Inactief BV", "KVK-nummer": "33333333", "raw_legal_form": "BV", "raw_status": "inactief"}])
    included, excluded, review = exclude_sole_proprietorships(run)
    assert len(read_tsv(included)) == 1 and len(read_tsv(excluded)) == 1 and len(read_tsv(review)) == 1
    active, inactive, status_review = active_only(run)
    assert read_tsv(active) == [] and len(read_tsv(inactive)) == 1 and read_tsv(status_review) == []
    unresolved = _register(run, "05", "kvk_unresolved", ["candidate_id", "reason"], [{"candidate_id": "x", "reason": "NOT_FOUND"}])
    with pytest.raises(HarvestError):
        export(run)
    assert export(run, allow_partial=True)
    unresolved.write_text("changed")
    with pytest.raises(HarvestError) as error:
        verify(run)
    assert error.value.exit_code == 7
    with pytest.raises(HarvestError):
        trace(run, "99999999")


def test_excel_safe_text(tmp_path: Path) -> None:
    path = tmp_path / "safe.xlsx"
    write_xlsx(path, ["Bedrijfsnaam", "KVK-nummer"], [{"Bedrijfsnaam": "=cmd", "KVK-nummer": "00123456"}])
    sheet = load_workbook(path)["Bedrijven"]
    assert sheet["A2"].value == "'=cmd" and sheet.freeze_panes == "A2"


def test_export_delivers_every_active_company_without_ranking_or_cap(run) -> None:
    _register(run, "07", "active", ["Bedrijfsnaam", "KVK-nummer"], [
        {"Bedrijfsnaam": "Zulu B.V.", "KVK-nummer": "33333333"},
        {"Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "11111111"},
        {"Bedrijfsnaam": "Midden B.V.", "KVK-nummer": "22222222"},
    ])
    paths = export(run)
    assert [row["Bedrijfsnaam"] for row in read_tsv(paths[0])] == [
        "Alpha B.V.", "Midden B.V.", "Zulu B.V.",
    ]
    assert read_tsv(paths[-2]) == []
    manifest = json.loads(paths[-1].read_text())
    assert manifest["selection_policy"] == "ALL_ACTIVE" and manifest["active_rows"] == 3


def test_light_export_rejects_wrong_source_number(run) -> None:
    _register(run, "03", "candidates", ["candidate_id", "source_kvk_hint", "website", "sector"], [
        {"candidate_id": "one", "source_kvk_hint": "11111111", "website": "https://wrong.example", "sector": "Bouw"},
    ])
    _register(run, "07", "active", ["candidate_id", "Bedrijfsnaam", "KVK-nummer"], [
        {"candidate_id": "one", "Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "22222222"},
    ])
    with pytest.raises(HarvestError, match="broncontext wijkt af"):
        export(run)


def test_light_export_rejects_unreviewed_response_metadata(run) -> None:
    _register(run, "07", "active", ["Bedrijfsnaam", "KVK-nummer", "response_json"], [
        {"Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "01234567",
         "response_json": json.dumps([{"naam": "Alpha B.V.", "metadata": {"internal": "secret"}}])},
    ])
    with pytest.raises(HarvestError, match="onbekend KVK-responsveld"):
        export(run)


def test_light_export_recovers_source_fields_from_merged_payloads(run) -> None:
    _register(run, "03", "candidates", ["candidate_id", "source_kvk_hint", "website", "sector", "source_payloads_json"], [
        {"candidate_id": "one", "source_kvk_hint": "01234567", "website": "", "sector": "",
         "source_payloads_json": json.dumps([{}, {"website": "https://alpha.example", "sector": "Bouw"}])},
    ])
    _register(run, "07", "active", ["candidate_id", "Bedrijfsnaam", "KVK-nummer"], [
        {"candidate_id": "one", "Bedrijfsnaam": "Alpha B.V.", "KVK-nummer": "01234567"},
    ])
    light = read_tsv(export(run)[2])[0]
    assert light["Website (bron)"] == "https://alpha.example"
    assert light["Sector (bron)"] == "Bouw"


def test_peak_memory_has_explicit_windows_fallback(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("company_harvest.workflow.sys.platform", "win32")
    assert _peak_memory() == (None, "UNAVAILABLE_ON_PLATFORM")


def test_partition_overlap_and_atomic_export_failure(run, monkeypatch: pytest.MonkeyPatch) -> None:
    assert _overlap([[{"KVK-nummer": "01234567"}], [{"KVK-nummer": "01234567"}]])
    _register(run, "07", "active", ["Bedrijfsnaam", "KVK-nummer"], [{"Bedrijfsnaam": "Alpha", "KVK-nummer": "01234567"}])
    monkeypatch.setattr("company_harvest.workflow.write_xlsx", lambda *args: (_ for _ in ()).throw(OSError("fault")))
    with pytest.raises(OSError):
        export(run)
    assert not list((run.path / "artifacts").glob("*_08_delivery_outputset"))
    assert not list((run.path / "artifacts").glob(".*_08_delivery_outputset.tmp"))


def test_audit_requires_terminal_artifacts(run) -> None:
    _register(run, "03", "candidates", ["candidate_id"], [{"candidate_id": "one"}])
    run.update_status("IN_PROGRESS", "04")
    with pytest.raises(HarvestError):
        verify(run)


def test_outcome_report_closes_empty_duplicate_conflict_and_review_fixtures(run) -> None:
    headers = [
        "original_name", "source_kvk_hint", "source_id", "source_registration_raw", "source_row",
        "source_url", "registration_validation_status", "source_legal_form", "source_status",
    ]
    raw = [
        {"original_name": "Alpha BV", "source_kvk_hint": "01234567", "source_id": "a", "source_registration_raw": "01234567", "source_row": "1", "source_url": "a", "registration_validation_status": "VALID"},
        {"original_name": " Alpha BV ", "source_kvk_hint": "01234567", "source_id": "b", "source_registration_raw": "01234567", "source_row": "1", "source_url": "b", "registration_validation_status": "VALID"},
        {"original_name": "Naamgenoot", "source_kvk_hint": "", "source_id": "a", "source_registration_raw": "", "source_row": "2", "source_url": "a", "registration_validation_status": "MISSING"},
        {"original_name": "Naamgenoot", "source_kvk_hint": "", "source_id": "b", "source_registration_raw": "", "source_row": "2", "source_url": "b", "registration_validation_status": "MISSING"},
        {"original_name": "Ongeldig", "source_kvk_hint": "", "source_id": "a", "source_registration_raw": "abc", "source_row": "3", "source_url": "a", "registration_validation_status": "INVALID"},
        {"original_name": "Ongeldig", "source_kvk_hint": "", "source_id": "b", "source_registration_raw": "xyz", "source_row": "3", "source_url": "b", "registration_validation_status": "INVALID"},
        {"original_name": "Conflict", "source_kvk_hint": "11111111", "source_id": "a", "source_registration_raw": "11111111", "source_row": "4", "source_url": "a", "registration_validation_status": "VALID"},
        {"original_name": "Conflict", "source_kvk_hint": "22222222", "source_id": "b", "source_registration_raw": "22222222", "source_row": "4", "source_url": "b", "registration_validation_status": "VALID"},
    ]
    _register(run, "01", "sources_inventory", ["source_id", "source_family"], [
        {"source_id": "a", "source_family": "family-a"},
        {"source_id": "b", "source_family": "family-b"},
    ])
    _register(run, "02", "source_a", headers, raw[::2])
    _register(run, "02", "source_b", headers, raw[1::2])
    candidates_path, decisions_path, conflicts_path = merge_candidates(run)
    assert len(read_tsv(candidates_path)) == 5
    assert len(read_tsv(decisions_path)) == 5
    assert len(read_tsv(conflicts_path)) == 2
    assert sum(row["original_name"] == "Naamgenoot" for row in read_tsv(candidates_path)) == 2
    metrics = outcome_metrics(run)
    assert metrics["counts"] == {
        "raw_records": 8,
        "valid_registration_numbers": 4,
        "without_direct_registration_number": 4,
        "missing_registration_numbers": 2,
        "invalid_registration_numbers": 2,
        "unique_candidates_before_deduplication": 7,
        "unique_candidates_after_deduplication": 5,
        "identical_merges": 1,
        "conflict_records": 2,
        "review_case_records": 6,
        "cross_source_candidates": 1,
    }
    assert metrics["count_closure"]["status"] == "PARTIAL_CLOSED"
    assert metrics["count_closure"]["transitions"]["raw_to_dedup"] == {
        "status": "CLOSED", "input": 8, "output": 8, "delta": 0,
    }
    assert metrics["source_diversity"]["source_family_count"] == 2
    assert metrics["source_diversity"]["measured_candidate_overlap_by_family_pair"] == {
        "family-a|family-b": 1,
    }
    assert metrics["source_diversity"][
        "measured_valid_registration_overlap_by_source_pair"
    ] == {"a|b": 1}
    assert metrics["resources"]["storage_growth_status"] == "MEASURED"
    empty_run = initialize_run(run.path.parent, 1)
    empty = outcome_metrics(empty_run)
    assert empty["counts"]["raw_records"] == 0
    assert empty["count_closure"]["status"] == "NOT_AVAILABLE"


def test_merge_and_metrics_use_only_latest_complete_artifact_per_source(run) -> None:
    headers = ["original_name", "source_kvk_hint", "source_id", "source_row", "source_url"]
    _register(run, "02", "source_a", headers, [{
        "original_name": "Oud BV", "source_kvk_hint": "", "source_id": "a",
        "source_row": "1", "source_url": "old",
    }])
    _register(run, "02", "source_a", headers, [{
        "original_name": "Nieuw BV", "source_kvk_hint": "", "source_id": "a",
        "source_row": "1", "source_url": "new",
    }])
    candidates, _, _ = merge_candidates(run)
    assert [row["original_name"] for row in read_tsv(candidates)] == ["Nieuw BV"]
    metrics = outcome_metrics(run)
    assert metrics["counts"]["raw_records"] == 1
    assert metrics["count_closure"]["transitions"]["raw_to_dedup"] == {
        "status": "CLOSED", "input": 1, "output": 1, "delta": 0,
    }
