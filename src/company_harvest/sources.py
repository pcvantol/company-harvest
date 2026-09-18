"""Begrensde publieke broncatalogus en extractie."""

from __future__ import annotations

import csv
import html
import io
import json
import re
import shutil
import time
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook

from company_harvest.core import (
    HTTP_USER_AGENT,
    HarvestError,
    Run,
    atomic_write,
    normalize_name,
    read_tsv,
    timestamp,
    validate_kvk,
    write_tsv,
)

MAX_RESPONSE_BYTES = 20 * 1024 * 1024
MAX_SOURCE_PAGES = 100
WIKIDATA_PAGE_SIZE = 100
DEFAULT_WIKIDATA_MEASUREMENT_LIMIT = 200
CAPABILITY_REPORT_SCHEMA_VERSION = 1
ALLOWED_HOSTS = {"ind.nl", "www.wikidata.org", "query.wikidata.org"}
DIRECT_COLLECT_SOURCE_IDS = {"ind_arbeid", "wikidata_nl_companies"}
CAPABILITY_SOURCE_IDS = ("ind_arbeid", "wikidata_nl_companies")


@dataclass(frozen=True)
class Source:
    source_id: str
    name: str
    owner: str
    source_family: str
    url: str
    parser: str
    access_mode: str
    terms_url: str
    refresh_info: str
    has_registration_number: bool
    registration_number_type: str
    provides_legal_form: bool
    provides_status: bool
    inclusion_reason: str
    provenance_quality_status: str
    bias: str


CATALOG = (
    Source(
        "ind_arbeid",
        "Openbaar register Arbeid",
        "IND",
        "overheidsregister",
        "https://ind.nl/nl/openbaar-register-erkende-referenten/openbaar-register-arbeid",
        "ind_html_table_v1",
        "html",
        "https://ind.nl/nl/proclaimer",
        "Maandelijks volgens de bronpagina; peildatum wordt per live meting vastgelegd.",
        True,
        "KVK",
        False,
        False,
        "Publiek Nederlands register met directe organisatie-identifiers.",
        "PRIMARY_OWNER_SOURCE",
        "Erkende referenten voor arbeid; geen sectorbrede populatie.",
    ),
    Source(
        "wikidata_nl_companies",
        "Wikidata Nederlandse organisaties",
        "Wikimedia community",
        "open-kennisbank",
        "https://query.wikidata.org/sparql",
        "wikidata_sparql_v1",
        "api",
        "https://www.wikidata.org/wiki/Wikidata:Data_access",
        "Doorlopend bewerkbaar; actualiteit verschilt per item.",
        True,
        "KVK",
        False,
        False,
        "Onafhankelijke kennisbron voor aanvullende dekking en overlapmeting.",
        "COMMUNITY_CURATED",
        "Vrijwillig samengestelde kennisbank; dekking en actualiteit variëren.",
    ),
    Source(
        "gleif_golden_copy",
        "GLEIF Level 1 Golden Copy",
        "Global Legal Entity Identifier Foundation",
        "wereldwijd-entiteitenregister",
        "https://goldencopy.gleif.org/api/v2/golden-copies/publishes/lei2/latest.csv",
        "gleif_golden_copy_csv_v1",
        "bulk",
        "https://www.gleif.org/en/meta/lei-data-terms-of-use",
        "Drie Golden Copy-publicaties per dag; de evidence-snapshot bepaalt de gebruikte versie.",
        True,
        "KVK via registratieautoriteit RA000463",
        True,
        True,
        "Brede primaire LEI-bron met herleidbare Nederlandse registratie-identifiers en bronvelden.",
        "PRIMARY_AGGREGATED_REGISTER",
        "Alleen entiteiten met een LEI; geen volledige populatie van Nederlandse ondernemingen.",
    ),
)

SOURCE_HEADERS = [
    "catalog_schema_version", "source_id", "name", "owner", "source_family", "url",
    "parser", "access_mode", "discovered_at", "terms_url", "refresh_info", "status",
    "live_measurement_status", "measured_count", "has_registration_number",
    "registration_number_type", "provides_legal_form", "provides_status",
    "estimated_overlap", "inclusion_reason", "provenance_quality_status",
    "candidate_layer", "bias",
]
RAW_HEADERS = [
    "original_name", "country", "nl_evidence", "employees_raw", "employees_date",
    "employees_scope", "website", "sector", "source_kvk_hint", "source_id", "source_url",
    "fetched_at", "source_row", "evidence_reference", "source_registration_raw",
    "registration_validation_status", "source_legal_form", "source_status", "candidate_layer",
]

CATALOG_SCHEMA_VERSION = "2"


def _catalog_row(source: Source, discovered_at: str) -> dict[str, str]:
    return {
        "catalog_schema_version": CATALOG_SCHEMA_VERSION,
        "source_id": source.source_id,
        "name": source.name,
        "owner": source.owner,
        "source_family": source.source_family,
        "url": source.url,
        "parser": source.parser,
        "access_mode": source.access_mode,
        "discovered_at": discovered_at,
        "terms_url": source.terms_url,
        "refresh_info": source.refresh_info,
        "status": "CONFIGURED_NOT_COLLECTED",
        "live_measurement_status": "NOT_MEASURED",
        "measured_count": "",
        "has_registration_number": str(source.has_registration_number).lower(),
        "registration_number_type": source.registration_number_type,
        "provides_legal_form": str(source.provides_legal_form).lower(),
        "provides_status": str(source.provides_status).lower(),
        "estimated_overlap": "",
        "inclusion_reason": source.inclusion_reason,
        "provenance_quality_status": source.provenance_quality_status,
        "candidate_layer": "raw",
        "bias": source.bias,
    }


def _normalize_catalog_row(row: dict[str, str]) -> dict[str, str]:
    known = next((source for source in CATALOG if source.source_id == row.get("source_id")), None)
    if known:
        normalized = _catalog_row(known, row.get("discovered_at", ""))
    else:
        normalized = {header: "" for header in SOURCE_HEADERS}
        normalized.update(
            {
                "catalog_schema_version": CATALOG_SCHEMA_VERSION,
                "source_family": "unclassified",
                "access_mode": "manual-import",
                "live_measurement_status": "NOT_MEASURED",
                "has_registration_number": "unknown",
                "provides_legal_form": "unknown",
                "provides_status": "unknown",
                "provenance_quality_status": "UNCLASSIFIED",
                "candidate_layer": "raw",
            }
        )
    for header in SOURCE_HEADERS:
        if header in row:
            normalized[header] = row[header]
    normalized["catalog_schema_version"] = CATALOG_SCHEMA_VERSION
    if not normalized["live_measurement_status"]:
        normalized["live_measurement_status"] = (
            "MEASURED" if normalized.get("measured_count") else "NOT_MEASURED"
        )
    return normalized


def discover(run: Run) -> tuple[Path, Path]:
    now = datetime.now(UTC).isoformat()
    rows = [_catalog_row(source, now) for source in CATALOG]
    csv_path = run.artifact_path("01", "sources_inventory", "csv")
    md_path = run.artifact_path("01", "sources_report", "md")
    write_tsv(csv_path, SOURCE_HEADERS, rows)
    report = "# Bronneninventaris\n\n" + "\n".join(
        f"- **{row['name']}** (`{row['source_id']}`): {row['status']}; "
        f"familie `{row['source_family']}`, toegang `{row['access_mode']}`, "
        f"registratie-ID `{row['registration_number_type'] or 'geen/unknown'}`. Bias: {row['bias']}"
        for row in rows
    ) + "\n\nAantallen blijven onbekend totdat de bron werkelijk is verzameld.\n"
    md_path.write_text(report, encoding="utf-8")
    run.register_artifact(csv_path, "01", "sources_inventory")
    run.register_artifact(md_path, "01", "sources_report")
    run.log("INFO", "sources_discovered", count=len(rows))
    run.update_status("IN_PROGRESS", "01")
    return csv_path, md_path


def _bounded_get(client: httpx.Client, url: str) -> httpx.Response:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in ALLOWED_HOSTS:
        raise HarvestError(f"bronhost niet toegestaan: {parsed.hostname}")
    response = client.get(url)
    if len(response.content) > MAX_RESPONSE_BYTES:
        raise HarvestError("bronresponse overschrijdt de maximale grootte")
    if response.url.host not in ALLOWED_HOSTS:
        raise HarvestError("bronredirect naar niet-toegestane host")
    return response


def _parse_ind_html_with_stats(
    content: str, source_url: str, limit: int | None = None
) -> tuple[list[dict[str, str]], int, int]:
    """Parse IND-tabelrijen en tel kandidaatrijen plus strikte afwijzingen."""
    table_rows = re.findall(r"<tr\b[^>]*>(.*?)</tr>", content, flags=re.I | re.S)
    rows: list[dict[str, str]] = []
    candidates = 0
    rejected = 0
    if table_rows:
        for block in table_rows:
            cells = re.findall(
                r"<(?:th|td)\b[^>]*>(.*?)</(?:th|td)>", block, flags=re.I | re.S
            )
            if len(cells) < 2:
                continue
            name = " ".join(html.unescape(re.sub(r"<[^>]+>", " ", cells[-2])).split())
            registration = " ".join(
                html.unescape(re.sub(r"<[^>]+>", " ", cells[-1])).split()
            )
            if not name or not registration or "kvk" in registration.casefold():
                continue
            candidates += 1
            try:
                kvk = validate_kvk(registration)
            except ValueError:
                rejected += 1
                continue
            rows.append(_raw_row(name, kvk, "ind_arbeid", source_url, str(candidates)))
            if limit and len(rows) >= limit:
                break
        return rows, candidates, rejected

    text = re.sub(r"<[^>]+>", "\n", content)
    text = html.unescape(text)
    lines = [" ".join(line.split()) for line in text.splitlines() if line.strip()]
    for index, line in enumerate(lines):
        inline = re.search(r"(.+?)\s+([0-9]{8})$", line)
        if inline:
            name, number = inline.groups()
        elif re.fullmatch(r"[0-9]{8}", line) and index:
            name, number = lines[index - 1], line
        else:
            continue
        candidates += 1
        try:
            kvk = validate_kvk(number)
        except ValueError:
            rejected += 1
            continue
        rows.append(_raw_row(name, kvk, "ind_arbeid", source_url, str(index + 1)))
        if limit and len(rows) >= limit:
            break
    return rows, candidates, rejected


def parse_ind_html(content: str, source_url: str, limit: int | None = None) -> list[dict[str, str]]:
    rows, _, _ = _parse_ind_html_with_stats(content, source_url, limit)
    if not rows:
        raise HarvestError("IND-parser vond onverwacht nul records", 5)
    return rows


def _ind_source_date(content: str) -> str | None:
    text = html.unescape(re.sub(r"<[^>]+>", " ", content))
    text = " ".join(text.split())
    match = re.search(r"Het overzicht is bijgewerkt op\s+([^.<]+)", text, flags=re.I)
    return match.group(1).strip() if match else None


def _raw_row(
    name: str,
    kvk: str,
    source_id: str,
    url: str,
    row: str,
    registration_raw: str | None = None,
    registration_status: str | None = None,
    nl_evidence: str = "Publiek Nederlands register",
) -> dict[str, str]:
    clean_name = re.sub(r'""([^\n]+?)""', r'"\1"', name.strip())
    return {
        "original_name": clean_name, "country": "Nederland",
        "nl_evidence": nl_evidence, "employees_raw": "",
        "employees_date": "", "employees_scope": "", "website": "", "sector": "",
        "source_kvk_hint": kvk, "source_id": source_id, "source_url": url,
        "fetched_at": datetime.now(UTC).isoformat(), "source_row": row,
        "evidence_reference": url,
        "source_registration_raw": kvk if registration_raw is None else registration_raw,
        "registration_validation_status": registration_status or ("VALID" if kvk else "MISSING"),
        "source_legal_form": "", "source_status": "", "candidate_layer": "raw",
    }


def parse_wikidata(payload: dict[str, object], source_url: str, limit: int | None = None) -> list[dict[str, str]]:
    results = payload.get("results")
    bindings = results.get("bindings", []) if isinstance(results, dict) else []
    rows: list[dict[str, str]] = []
    for index, binding in enumerate(bindings):
        if not isinstance(binding, dict):
            continue
        label = binding.get("orgLabel", {})
        kvk_cell = binding.get("kvk", {})
        name = label.get("value", "") if isinstance(label, dict) else ""
        kvk = kvk_cell.get("value", "") if isinstance(kvk_cell, dict) else ""
        try:
            number = validate_kvk(kvk)
            validation_status = "VALID"
        except ValueError:
            number = ""
            validation_status = "INVALID" if str(kvk).strip() else "MISSING"
        if str(name).strip():
            rows.append(
                _raw_row(
                    str(name),
                    number,
                    "wikidata_nl_companies",
                    source_url,
                    str(index + 1),
                    str(kvk),
                    validation_status,
                    "Wikidata-eigenschap P3220 (KvK company ID)",
                )
            )
        if limit and len(rows) >= limit:
            break
    return rows


def _enabled(only: Iterable[str], skip: Iterable[str]) -> list[Source]:
    only_set, skip_set = set(only), set(skip)
    known = {source.source_id for source in CATALOG}
    unknown = (only_set | skip_set) - known
    if unknown:
        raise HarvestError(f"onbekende source-id(s): {', '.join(sorted(unknown))}")
    if "gleif_golden_copy" in only_set:
        raise HarvestError("gebruik 'sources gleif' voor de begrensde GLEIF-bulkadapter")
    return [
        source
        for source in CATALOG
        if source.source_id in DIRECT_COLLECT_SOURCE_IDS
        and (not only_set or source.source_id in only_set)
        and source.source_id not in skip_set
    ]


def _record_source_observation(
    run: Run,
    observations: dict[str, object],
    source_id: str,
    started: float,
    http_statuses: list[int],
    response_count: int,
    response_bytes: int,
    raw_response_records: int,
    parser_rejections: int,
    limit: int | None,
    collection_complete: bool,
    rate_limit_headers: dict[str, str],
    source_data_date: str | None,
) -> None:
    observations[source_id] = {
        "observed_at": datetime.now(UTC).isoformat(),
        "duration_seconds": round(time.monotonic() - started, 6),
        "http_statuses": http_statuses,
        "response_count": response_count,
        "response_bytes": response_bytes,
        "raw_response_records": raw_response_records,
        "parser_rejections": parser_rejections,
        "limit": limit,
        "collection_complete": collection_complete,
        "rate_limit_headers": rate_limit_headers,
        "user_agent": HTTP_USER_AGENT,
        "source_data_date": source_data_date,
    }
    run.record_config("source_observations", observations)


def collect(run: Run, only: Iterable[str] = (), skip: Iterable[str] = (), limit: int | None = None, refresh: bool = False) -> list[Path]:
    outputs: list[Path] = []
    enabled = _enabled(only, skip)
    will_fetch = any(refresh or not run.latest_artifact("02", f"source_{source.source_id}") for source in enabled)
    if will_fetch:
        run.invalidate_from(3, "source_collection_changed")
    timeout = httpx.Timeout(20, connect=10, read=20, write=10, pool=10)
    headers = {"User-Agent": HTTP_USER_AGENT}
    observations = dict(run.metadata().get("runtime_config", {}).get("source_observations", {}))
    with httpx.Client(timeout=timeout, follow_redirects=True, max_redirects=3, headers=headers) as client:
        for source in enabled:
            prior = run.latest_artifact("02", f"source_{source.source_id}")
            if prior and not refresh:
                run.log("INFO", "source_collect_reused", source_id=source.source_id, path=str(prior.relative_to(run.path)))
                outputs.append(prior)
                continue
            run.log("INFO", "source_collect_started", source_id=source.source_id)
            started = time.monotonic()
            snapshots: list[Path] = []
            response_bytes = 0
            response_count = 0
            raw_response_records = 0
            parser_rejections = 0
            http_statuses: list[int] = []
            rate_limit_headers: dict[str, str] = {}
            collection_complete = False
            source_data_date: str | None = None

            if source.source_id == "ind_arbeid":
                response = _bounded_get(client, source.url)
                response_bytes = len(response.content)
                response_count = 1
                http_statuses.append(int(getattr(response, "status_code", 200)))
                rate_limit_headers.update(_rate_limit_headers(response))
                snapshot = run.path / "evidence" / f"{timestamp()}_02_ind_arbeid_source.html"
                snapshot.write_bytes(response.content)
                snapshots.append(snapshot)
                run.register_artifact(snapshot, "02", f"evidence_{source.source_id}")
                _record_source_observation(
                    run, observations, source.source_id, started, http_statuses,
                    response_count, response_bytes, raw_response_records,
                    parser_rejections, limit, collection_complete, rate_limit_headers,
                    source_data_date,
                )
                response.raise_for_status()
                rows, raw_response_records, parser_rejections = _parse_ind_html_with_stats(
                    response.text, str(response.url), limit
                )
                source_data_date = _ind_source_date(response.text)
                _record_source_observation(
                    run, observations, source.source_id, started, http_statuses,
                    response_count, response_bytes, raw_response_records,
                    parser_rejections, limit, collection_complete, rate_limit_headers,
                    source_data_date,
                )
                if not rows:
                    raise HarvestError("IND-parser vond onverwacht nul records", 5)
                collection_complete = limit is None
            else:
                rows = []
                for page in range(MAX_SOURCE_PAGES):
                    remaining = None if limit is None else limit - len(rows)
                    if remaining is not None and remaining <= 0:
                        break
                    page_size = min(WIKIDATA_PAGE_SIZE, remaining) if remaining is not None else WIKIDATA_PAGE_SIZE
                    query = f"SELECT ?org ?orgLabel ?kvk WHERE {{?org wdt:P3220 ?kvk. SERVICE wikibase:label {{ bd:serviceParam wikibase:language 'nl,en'. }} }} ORDER BY ?org LIMIT {page_size} OFFSET {page * WIKIDATA_PAGE_SIZE}"
                    response = client.get(source.url, params={"query": query, "format": "json"})
                    response_count += 1
                    response_bytes += len(response.content)
                    http_statuses.append(int(getattr(response, "status_code", 200)))
                    rate_limit_headers.update(_rate_limit_headers(response))
                    if response.url.host not in ALLOWED_HOSTS:
                        raise HarvestError("Wikidata-redirect naar niet-toegestane host")
                    if len(response.content) > MAX_RESPONSE_BYTES:
                        raise HarvestError("Wikidata-response overschrijdt de maximale grootte")
                    snapshot = run.path / "evidence" / f"{timestamp()}_02_wikidata_source_page_{page + 1}.json"
                    snapshot.write_bytes(response.content)
                    snapshots.append(snapshot)
                    run.register_artifact(snapshot, "02", f"evidence_{source.source_id}")
                    _record_source_observation(
                        run, observations, source.source_id, started, http_statuses,
                        response_count, response_bytes, raw_response_records,
                        parser_rejections, limit, collection_complete, rate_limit_headers,
                        source_data_date,
                    )
                    response.raise_for_status()
                    try:
                        payload = response.json()
                    except ValueError as exc:
                        raise HarvestError("Wikidata-response bevat ongeldige JSON") from exc
                    results = payload.get("results") if isinstance(payload, dict) else None
                    bindings = results.get("bindings", []) if isinstance(results, dict) else []
                    if not isinstance(bindings, list):
                        raise HarvestError("Wikidata-responseschema is ongeldig")
                    raw_response_records += len(bindings)
                    page_rows = parse_wikidata(payload, str(response.url), page_size)
                    for row in page_rows:
                        row["source_row"] = str(page * WIKIDATA_PAGE_SIZE + int(row["source_row"]))
                    rows.extend(page_rows)
                    if len(bindings) < page_size:
                        collection_complete = True
                        break
                else:
                    raise HarvestError("Wikidata-paginering bereikte de veiligheidslimiet; resultaat is niet aantoonbaar compleet")
            path = run.artifact_path("02", f"source_{source.source_id}_companies", "csv")
            rows.sort(key=lambda row: (row["original_name"].casefold(), row["source_kvk_hint"]))
            write_tsv(path, RAW_HEADERS, rows)
            run.register_artifact(path, "02", f"source_{source.source_id}")
            parser_rejections = (
                parser_rejections
                if source.source_id == "ind_arbeid"
                else max(raw_response_records - len(rows), 0)
            )
            _record_source_observation(
                run, observations, source.source_id, started, http_statuses,
                response_count, response_bytes, raw_response_records,
                parser_rejections, limit, collection_complete, rate_limit_headers,
                source_data_date,
            )
            run.log("INFO", "source_collect_completed", source_id=source.source_id, count=len(rows), evidence_sha256=[_hash(snapshot) for snapshot in snapshots], refreshed=refresh)
            outputs.append(path)
    _update_inventory(run, {path.name.split("_source_", 1)[1].rsplit("_companies", 1)[0]: len(read_tsv(path)) for path in outputs})
    run.record_config("sources_collect", {"only": sorted(only), "skip": sorted(skip), "limit": limit, "refresh": refresh})
    run.update_status("IN_PROGRESS", "02")
    return outputs


def _rate_limit_headers(response: httpx.Response) -> dict[str, str]:
    headers = getattr(response, "headers", {})
    names = ("retry-after", "x-ratelimit-limit", "x-ratelimit-remaining", "x-ratelimit-reset")
    return {name: str(headers[name]) for name in names if name in headers}


def _measurement_error(exc: Exception) -> tuple[str, dict[str, object]]:
    if isinstance(exc, httpx.HTTPStatusError):
        status_code = exc.response.status_code
        status = "BLOCKED" if status_code in {403, 429} else "FAILED"
        return status, {
            "error_type": type(exc).__name__,
            "http_status": status_code,
            "retry_after": exc.response.headers.get("retry-after"),
        }
    return "FAILED", {"error_type": type(exc).__name__, "message": str(exc)}


def _source_capability(
    source: Source,
    rows: list[dict[str, str]],
    observation: dict[str, object],
    measurement_status: str,
    error: dict[str, object] | None = None,
) -> dict[str, object]:
    valid_numbers = [
        row.get("source_kvk_hint", "")
        for row in rows
        if row.get("registration_validation_status") == "VALID"
        or re.fullmatch(r"[0-9]{8}", row.get("source_kvk_hint", ""))
    ]
    missing = sum(
        row.get("registration_validation_status") == "MISSING"
        or (not row.get("source_kvk_hint") and not row.get("source_registration_raw"))
        for row in rows
    )
    invalid = sum(
        row.get("registration_validation_status") == "INVALID"
        or (not row.get("source_kvk_hint") and bool(row.get("source_registration_raw")))
        for row in rows
    )
    normalized_names = [normalize_name(row.get("original_name", "")) for row in rows]
    parser_rejections_value = observation.get("parser_rejections", 0)
    parser_rejections = (
        parser_rejections_value if isinstance(parser_rejections_value, int) else 0
    )
    return {
        "source_id": source.source_id,
        "source_family": source.source_family,
        "measurement_status": measurement_status,
        "measurement_scope": (
            "FULL_SOURCE_PAGE" if source.source_id == "ind_arbeid" else "BOUNDED_SAMPLE"
        ),
        "collection_complete": bool(observation.get("collection_complete", False)),
        "source_data_date": observation.get("source_data_date"),
        "configured_limit": observation.get("limit"),
        "raw_records": len(rows),
        "valid_registration_numbers": len(valid_numbers),
        "unique_valid_registration_numbers": len(set(valid_numbers)),
        "missing_registration_numbers": missing,
        "invalid_registration_numbers": invalid,
        "duplicate_registration_records": len(valid_numbers) - len(set(valid_numbers)),
        "duplicate_name_records": len(normalized_names) - len(set(normalized_names)),
        "parser_rejections": parser_rejections,
        "count_closure": (
            "NOT_AVAILABLE"
            if measurement_status != "LIVE_MEASURED"
            else "CLOSED"
            if len(rows) == len(valid_numbers) + missing + invalid
            else "OPEN"
        ),
        "duration_seconds": observation.get("duration_seconds"),
        "response_count": observation.get("response_count"),
        "response_bytes": observation.get("response_bytes"),
        "http_statuses": observation.get("http_statuses", []),
        "rate_limit_headers": observation.get("rate_limit_headers", {}),
        "rate_limit_observation": (
            "HEADERS_OBSERVED" if observation.get("rate_limit_headers") else "NO_HEADERS_OBSERVED"
        ),
        "terms_url": source.terms_url,
        "refresh_info": source.refresh_info,
        "bias": source.bias,
        "provides_legal_form": source.provides_legal_form,
        "provides_status": source.provides_status,
        "error": error,
    }


def measure_sources(
    run: Run,
    wikidata_limit: int = DEFAULT_WIKIDATA_MEASUREMENT_LIMIT,
) -> tuple[Path, Path]:
    """Meet bestaande bronnen live en schrijft ook bij blokkades een terminal rapport."""
    if wikidata_limit < 1 or wikidata_limit > 1000:
        raise HarvestError("wikidata-meetlimiet moet tussen 1 en 1000 liggen")
    if not run.latest_artifact("01", "sources_inventory"):
        discover(run)
    attempts: dict[str, tuple[str, dict[str, object] | None]] = {}
    source_paths: dict[str, Path] = {}
    capability_sources = [
        source for source in CATALOG if source.source_id in CAPABILITY_SOURCE_IDS
    ]
    for source in capability_sources:
        limit = None if source.source_id == "ind_arbeid" else wikidata_limit
        observations = dict(
            run.metadata().get("runtime_config", {}).get("source_observations", {})
        )
        observations.pop(source.source_id, None)
        run.record_config("source_observations", observations)
        try:
            paths = collect(run, only=[source.source_id], limit=limit, refresh=True)
            source_paths[source.source_id] = paths[0]
            attempts[source.source_id] = ("LIVE_MEASURED", None)
        except (HarvestError, httpx.HTTPError) as exc:
            status, error = _measurement_error(exc)
            attempts[source.source_id] = (status, error)
            run.log(
                "WARNING",
                "source_measurement_terminal_error",
                source_id=source.source_id,
                status=status,
                **error,
            )

    observations = run.metadata().get("runtime_config", {}).get("source_observations", {})
    capabilities: list[dict[str, object]] = []
    valid_by_source: dict[str, set[str]] = {}
    for source in capability_sources:
        path = source_paths.get(source.source_id)
        rows = read_tsv(path) if path else []
        observation = observations.get(source.source_id, {}) if isinstance(observations, dict) else {}
        if not isinstance(observation, dict):
            observation = {}
        if "limit" not in observation:
            observation["limit"] = None if source.source_id == "ind_arbeid" else wikidata_limit
        status, attempt_error = attempts[source.source_id]
        capability = _source_capability(
            source,
            rows,
            observation,
            status,
            attempt_error,
        )
        capabilities.append(capability)
        valid_by_source[source.source_id] = {
            row["source_kvk_hint"]
            for row in rows
            if re.fullmatch(r"[0-9]{8}", row.get("source_kvk_hint", ""))
        }

    left_id, right_id = CAPABILITY_SOURCE_IDS
    live_by_source = {
        source_id: attempts[source_id][0] == "LIVE_MEASURED" for source_id in valid_by_source
    }
    overlap_available = all(live_by_source.values())
    shared = valid_by_source[left_id] & valid_by_source[right_id]
    report: dict[str, object] = {
        "capability_report_schema_version": CAPABILITY_REPORT_SCHEMA_VERSION,
        "generated_at": datetime.now(UTC).isoformat(),
        "run_id": run.metadata()["run_id"],
        "http_user_agent": HTTP_USER_AGENT,
        "measurement_kind": "BOUNDED_LIVE_CAPABILITY",
        "sources": capabilities,
        "exact_registration_overlap": {
            "status": "MEASURED" if overlap_available else "NOT_AVAILABLE",
            "source_pair": f"{left_id}|{right_id}",
            "shared_valid_registration_numbers": len(shared) if overlap_available else None,
            "left_unique_valid_registration_numbers": len(valid_by_source[left_id]),
            "right_unique_valid_registration_numbers": len(valid_by_source[right_id]),
            "left_overlap_share": (
                len(shared) / len(valid_by_source[left_id])
                if overlap_available and valid_by_source[left_id]
                else None
            ),
            "right_overlap_share": (
                len(shared) / len(valid_by_source[right_id])
                if overlap_available and valid_by_source[right_id]
                else None
            ),
        },
        "all_sources_terminal": all(
            item["measurement_status"] in {"LIVE_MEASURED", "BLOCKED", "FAILED"}
            for item in capabilities
        ),
        "thresholds": {"status": "NOT_SET_PENDING_R3_FEASIBILITY"},
    }
    json_path = run.artifact_path("02", "source_capability_report", "json")
    atomic_write(json_path, json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    lines = [
        "# Broncapabilitymeting",
        "",
        f"- Run: `{report['run_id']}`",
        f"- Meetsoort: `{report['measurement_kind']}`",
        f"- User-Agent: `{HTTP_USER_AGENT}`",
        f"- Alle bronnen terminaal: **{str(report['all_sources_terminal']).lower()}**",
        "",
        "| Bron | Status | Scope | Compleet | Records | Geldig KVK | Uniek KVK | Duplicaat-ID | Parserafwijzingen |",
        "|---|---|---|---|---:|---:|---:|---:|---:|",
    ]
    lines.extend(
        f"| {item['source_id']} | {item['measurement_status']} | {item['measurement_scope']} | "
        f"{item['collection_complete']} | {item['raw_records']} | "
        f"{item['valid_registration_numbers']} | {item['unique_valid_registration_numbers']} | "
        f"{item['duplicate_registration_records']} | {item['parser_rejections']} |"
        for item in capabilities
    )
    overlap = report["exact_registration_overlap"]
    assert isinstance(overlap, dict)
    lines.extend(
        [
            "",
            "## Exacte KVK-overlap",
            "",
            f"- Bronpaar: `{overlap['source_pair']}`",
            f"- Status: `{overlap['status']}`",
            f"- Gedeelde geldige nummers: {overlap['shared_valid_registration_numbers']}",
            "",
            "Rechtsvorm en status zijn uitsluitend bronvelden en gelden niet als KVK-gevalideerd. "
            "Er zijn nog geen succesdrempels vastgesteld.",
            "",
        ]
    )
    md_path = run.artifact_path("02", "source_capability_report", "md")
    atomic_write(md_path, "\n".join(lines))
    run.register_artifact_set(
        [
            (json_path, "02", "source_capability_report", "COMPLETE"),
            (md_path, "02", "source_capability_report_md", "COMPLETE"),
        ]
    )
    _update_capability_inventory(run, capabilities, len(shared) if overlap_available else None)
    run.record_config(
        "sources_measure",
        {"wikidata_limit": wikidata_limit, "refresh": True, "report_schema": 1},
    )
    run.log("INFO", "source_capability_measurement_completed", terminal=True)
    return json_path, md_path


def _update_capability_inventory(
    run: Run, capabilities: list[dict[str, object]], shared_count: int | None
) -> None:
    rows = read_catalog(run)
    by_id = {str(item["source_id"]): item for item in capabilities}
    for row in rows:
        capability = by_id.get(row["source_id"])
        if not capability:
            continue
        status = str(capability["measurement_status"])
        row["status"] = "COLLECTED" if status == "LIVE_MEASURED" else status
        row["live_measurement_status"] = status
        row["measured_count"] = (
            str(capability["raw_records"]) if status == "LIVE_MEASURED" else ""
        )
        row["estimated_overlap"] = (
            f"exact_shared_kvk:{shared_count}" if shared_count is not None else "NOT_AVAILABLE"
        )
    updated = run.artifact_path("01", "sources_inventory_capability_measured", "csv")
    write_tsv(updated, SOURCE_HEADERS, rows)
    run.register_artifact(updated, "01", "sources_inventory")


def _hash(path: Path) -> str:
    import hashlib

    return hashlib.sha256(path.read_bytes()).hexdigest()


def list_sources(run: Run) -> list[dict[str, str]]:
    return read_catalog(run, migrate=True)


def read_catalog(run: Run, migrate: bool = False) -> list[dict[str, str]]:
    path = run.latest_artifact("01", "sources_inventory")
    if not path:
        raise HarvestError("voer eerst sources discover uit")
    rows = read_tsv(path)
    normalized = [_normalize_catalog_row(row) for row in rows]
    if migrate and rows and list(rows[0]) != SOURCE_HEADERS:
        migrated = run.artifact_path("01", "sources_inventory_schema_2", "csv")
        write_tsv(migrated, SOURCE_HEADERS, normalized)
        run.register_artifact(migrated, "01", "sources_inventory")
        run.log("INFO", "sources_inventory_migrated", from_schema="legacy", to_schema=2)
    return normalized


def _update_inventory(run: Run, counts: dict[str, int]) -> None:
    inventory = run.latest_artifact("01", "sources_inventory")
    if not inventory:
        return
    rows = [_normalize_catalog_row(row) for row in read_tsv(inventory)]
    for row in rows:
        if row["source_id"] in counts:
            row["status"] = "COLLECTED"
            row["live_measurement_status"] = "MEASURED"
            row["measured_count"] = str(counts[row["source_id"]])
    updated = run.artifact_path("01", "sources_inventory_measured", "csv")
    write_tsv(updated, SOURCE_HEADERS, rows)
    run.register_artifact(updated, "01", "sources_inventory")


def import_source(
    run: Run,
    input_path: Path,
    source_id: str,
    name_column: str,
    kvk_column: str | None = None,
    sheet: str | None = None,
) -> Path:
    """Importeer een expliciet gemapte lokale CSV/TSV/XLSX/HTML-tabel als bron."""
    if not re.fullmatch(r"[a-z0-9][a-z0-9_-]{1,63}", source_id):
        raise HarvestError("source-id moet 2-64 veilige kleine letters/cijfers bevatten")
    source = input_path.expanduser().resolve()
    if not source.is_file() or source.stat().st_size > MAX_RESPONSE_BYTES:
        raise HarvestError("importbron ontbreekt of overschrijdt de maximale grootte")
    headers, values = _read_tabular(source, sheet)
    if name_column not in headers or (kvk_column and kvk_column not in headers):
        raise HarvestError("expliciete naam- of opgegeven KVK-kolom ontbreekt")
    snapshot = run.path / "evidence" / f"{timestamp()}_02_{source_id}_input{source.suffix.lower()}"
    shutil.copyfile(source, snapshot)
    rows: list[dict[str, str]] = []
    for index, row in enumerate(values, 2):
        name = str(row.get(name_column) or "").strip()
        if not name:
            continue
        registration_raw = str(row.get(kvk_column) or "").strip() if kvk_column else ""
        try:
            number = validate_kvk(registration_raw)
            validation_status = "VALID"
        except ValueError:
            number = ""
            validation_status = "INVALID" if registration_raw else "MISSING"
        rows.append(
            _raw_row(
                name,
                number,
                source_id,
                snapshot.name,
                str(index),
                registration_raw,
                validation_status,
            )
        )
    if not rows:
        raise HarvestError("importadapter vond nul records met een bruikbare naam")
    run.invalidate_from(3, "source_import_changed")
    output = run.artifact_path("02", f"source_{source_id}_companies", "csv")
    write_tsv(output, RAW_HEADERS, rows)
    run.register_artifact(snapshot, "02", f"evidence_{source_id}")
    run.register_artifact(output, "02", f"source_{source_id}")
    _upsert_imported_source(run, source_id, source.suffix.lower(), kvk_column is not None, len(rows))
    run.log("INFO", "source_import_completed", source_id=source_id, count=len(rows), adapter=source.suffix.lower(), evidence_sha256=_hash(snapshot))
    run.record_config(f"source_import:{source_id}", {"sha256": _hash(snapshot), "adapter": source.suffix.lower(), "name_column": name_column, "kvk_column": kvk_column, "sheet": sheet})
    run.update_status("IN_PROGRESS", "02")
    return output


def _upsert_imported_source(
    run: Run, source_id: str, adapter: str, has_registration_number: bool, count: int
) -> None:
    inventory = run.latest_artifact("01", "sources_inventory")
    if inventory:
        rows = [_normalize_catalog_row(row) for row in read_tsv(inventory)]
    else:
        now = datetime.now(UTC).isoformat()
        rows = [_catalog_row(source, now) for source in CATALOG]
    rows = [row for row in rows if row["source_id"] != source_id]
    rows.append(
        {
            **{header: "" for header in SOURCE_HEADERS},
            "catalog_schema_version": CATALOG_SCHEMA_VERSION,
            "source_id": source_id,
            "name": source_id,
            "owner": "Lokale import",
            "source_family": "manual-import",
            "url": "lokale-evidence-snapshot",
            "parser": f"tabular{adapter}",
            "access_mode": "manual-import",
            "discovered_at": datetime.now(UTC).isoformat(),
            "refresh_info": "Per expliciete lokale import.",
            "status": "COLLECTED",
            "live_measurement_status": "MEASURED_LOCAL",
            "measured_count": str(count),
            "has_registration_number": str(has_registration_number).lower(),
            "registration_number_type": "KVK" if has_registration_number else "",
            "provides_legal_form": "false",
            "provides_status": "false",
            "inclusion_reason": "Expliciet door de gebruiker aangeleverde bron.",
            "provenance_quality_status": "USER_PROVIDED_UNVERIFIED",
            "candidate_layer": "raw",
            "bias": "Dekking en selectie volgen uit de aangeleverde bron.",
        }
    )
    updated = run.artifact_path("01", "sources_inventory_imported", "csv")
    write_tsv(updated, SOURCE_HEADERS, sorted(rows, key=lambda row: row["source_id"]))
    run.register_artifact(updated, "01", "sources_inventory")


def _read_tabular(path: Path, sheet: str | None) -> tuple[list[str], list[dict[str, object]]]:
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in workbook.sheetnames:
            raise HarvestError(f"werkblad ontbreekt: {sheet}")
        selected = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
        matrix = [[cell.value for cell in row] for row in selected.iter_rows()]
    elif path.suffix.lower() in {".html", ".htm"}:
        rows = re.findall(r"<tr\b[^>]*>(.*?)</tr>", path.read_text(encoding="utf-8"), flags=re.I | re.S)
        matrix = [[html.unescape(re.sub(r"<[^>]+>", "", cell)).strip() for cell in re.findall(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", row, flags=re.I | re.S)] for row in rows]
    else:
        text = path.read_text(encoding="utf-8-sig")
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters="\t,;").delimiter
        except csv.Error:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        if not reader.fieldnames:
            raise HarvestError("importheader ontbreekt")
        return list(reader.fieldnames), [dict(row) for row in reader]
    if not matrix:
        raise HarvestError("lege importbron")
    headers = [str(value or "").strip() for value in matrix[0]]
    if not all(headers) or len(headers) != len(set(headers)):
        raise HarvestError("importheader bevat lege of dubbele kolommen")
    return headers, [dict(zip(headers, row, strict=False)) for row in matrix[1:]]
