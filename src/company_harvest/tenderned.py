"""Publieke TenderNed-gunningen als lokale, controleerbare vijfde bron."""

from __future__ import annotations

import csv
import html
import json
import os
import re
import shutil
import time
import zipfile
from collections import Counter
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import cast
from urllib.parse import urljoin, urlparse

import httpx
from openpyxl import load_workbook

from company_harvest.core import (
    HTTP_USER_AGENT,
    HarvestError,
    Run,
    atomic_write,
    sha256,
    timestamp,
    validate_kvk,
)
from company_harvest.sources import RAW_HEADERS, SOURCE_HEADERS, read_catalog

SOURCE_ID = "tenderned_awards"
PAGE_URL = "https://www.tenderned.nl/cms/nl/aanbesteden-in-cijfers/datasets-aanbestedingen"
HOST = "www.tenderned.nl"
MAX_PAGE_BYTES = 2 * 1024 * 1024
MAX_FILE_BYTES = 160 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED = 900 * 1024 * 1024
MAX_REDIRECTS = 3
MAX_DOWNLOAD_SECONDS = 15 * 60
REJECTED_HEADERS = ["source_row", "original_name", "source_registration_raw", "country", "reason"]
REQUIRED_XLSX = {
    "ID publicatie", "Publicatiedatum", "Naam gegunde onderneming", "ON kvknummer",
    "ON land", "ON plaats", "URL TenderNed",
}


class _Links(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(html.unescape(href))


def _safe_url(url: str) -> None:
    parsed = urlparse(url)
    try:
        port = parsed.port
    except ValueError as exc:
        raise HarvestError("TenderNed-URL heeft een ongeldige poort") from exc
    if (parsed.scheme != "https" or parsed.hostname != HOST or port not in (None, 443)
            or parsed.username or parsed.password or parsed.fragment):
        raise HarvestError("TenderNed-URL buiten officiële HTTPS-host")


def _download(url: str, destination: Path, maximum: int) -> dict[str, object]:
    temporary = destination.with_name(f".{destination.name}.tmp")
    started = time.monotonic()
    downloaded = 0
    metadata: dict[str, object] = {}
    try:
        with httpx.Client(
            timeout=httpx.Timeout(120, connect=15), follow_redirects=False,
            headers={"User-Agent": HTTP_USER_AGENT},
        ) as client:
            current = url
            for redirects in range(MAX_REDIRECTS + 1):
                _safe_url(current)
                with client.stream("GET", current) as response:
                    _safe_url(str(response.url))
                    if 300 <= response.status_code < 400:
                        location = response.headers.get("location")
                        if not location or redirects == MAX_REDIRECTS:
                            raise HarvestError("TenderNed-download overschrijdt redirectgrens")
                        current = urljoin(str(response.url), location)
                        _safe_url(current)
                        continue
                    try:
                        length = int(response.headers["content-length"]) if "content-length" in response.headers else None
                    except ValueError as exc:
                        raise HarvestError("TenderNed-download heeft ongeldige Content-Length") from exc
                    if length is not None and (length < 1 or length > maximum):
                        raise HarvestError("TenderNed-download overschrijdt groottegate")
                    response.raise_for_status()
                    with temporary.open("wb") as handle:
                        for chunk in response.iter_bytes(1024 * 1024):
                            if time.monotonic() - started > MAX_DOWNLOAD_SECONDS:
                                raise HarvestError("TenderNed-download overschrijdt tijdgate")
                            downloaded += len(chunk)
                            if downloaded > maximum:
                                raise HarvestError("TenderNed-download overschrijdt groottegate")
                            handle.write(chunk)
                        handle.flush()
                        os.fsync(handle.fileno())
                    metadata = {
                        "url": str(response.url), "http_status": response.status_code,
                        "redirects": redirects, "content_length": length,
                        "last_modified": response.headers.get("last-modified"),
                        "downloaded_bytes": downloaded, "user_agent": HTTP_USER_AGENT,
                    }
                    break
    except httpx.HTTPStatusError as exc:
        raise HarvestError(f"TenderNed-download stopte met HTTP {exc.response.status_code}") from exc
    except httpx.HTTPError as exc:
        raise HarvestError(f"TenderNed-download mislukt: {type(exc).__name__}") from exc
    finally:
        if temporary.exists() and not metadata:
            temporary.unlink()
    if downloaded == 0:
        raise HarvestError("TenderNed-download is leeg")
    os.replace(temporary, destination)
    return metadata


def _select_downloads(page: str) -> tuple[str, str, int]:
    parser = _Links()
    parser.feed(page)
    xlsx: list[tuple[str, str]] = []
    json_files: list[tuple[str, str, str]] = []
    for href in parser.links:
        url = urljoin(PAGE_URL, href)
        try:
            _safe_url(url)
        except HarvestError:
            continue
        name = Path(urlparse(url).path).name
        excel_match = re.fullmatch(
            r"Dataset_Tenderned-compleet-2021-01-01-(\d{4}-\d{2}-\d{2})\.xlsx", name,
            re.IGNORECASE,
        )
        json_match = re.fullmatch(
            r"Dataset_Tenderned-(\d{4})-01-01-(\d{4}-\d{2}-\d{2})\.json", name,
            re.IGNORECASE,
        )
        if excel_match:
            xlsx.append((excel_match.group(1), url))
        if json_match:
            json_files.append((json_match.group(1), json_match.group(2), url))
    if not xlsx:
        raise HarvestError("TenderNed-pagina mist verwachte XLSX vanaf 2021")
    end_date, xlsx_url = max(xlsx)
    latest_year = int(end_date[:4])
    matches = [url for year, end, url in json_files if int(year) == latest_year and end == end_date]
    if len(matches) != 1:
        raise HarvestError("TenderNed-pagina mist unieke JSON voor de laatste XLSX-periode")
    return xlsx_url, matches[0], latest_year


def _local_copy(source: Path, destination: Path, maximum: int) -> None:
    source = source.expanduser().resolve()
    if not source.is_file() or not 0 < source.stat().st_size <= maximum:
        raise HarvestError("TenderNed-lokaal bronbestand ontbreekt of overschrijdt groottegate")
    if shutil.disk_usage(destination.parent).free < source.stat().st_size + 256 * 1024 * 1024:
        raise HarvestError("onvoldoende vrije ruimte voor TenderNed-evidence")
    temporary = destination.with_name(f".{destination.name}.tmp")
    try:
        shutil.copyfile(source, temporary)
        if sha256(source) != sha256(temporary):
            raise HarvestError("TenderNed-evidencekopie wijkt af")
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _safe_xlsx(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as bundle:
            entries = bundle.infolist()
            if not entries or sum(item.file_size for item in entries) > MAX_XLSX_UNCOMPRESSED:
                raise HarvestError("TenderNed-XLSX overschrijdt ongecomprimeerde groottegate")
            for item in entries:
                normalized_name = item.filename.replace("\\", "/")
                if (normalized_name.startswith("/") or ".." in normalized_name.split("/")
                        or item.flag_bits & 1 or item.file_size / max(item.compress_size, 1) > 30):
                    raise HarvestError("TenderNed-XLSX bevat onveilige ZIP-entry")
    except zipfile.BadZipFile as exc:
        raise HarvestError("TenderNed-XLSX is geen geldige ZIP") from exc


def _kvk(raw: object) -> tuple[str, str]:
    if raw is None or not str(raw).strip():
        return "", "MISSING"
    try:
        return validate_kvk(raw), "VALID"
    except ValueError:
        return "", "INVALID"


def _candidate(
    *, name: str, raw_id: object, city: str, source_row: str, source_url: str,
    evidence_name: str, fetched_at: str,
) -> tuple[dict[str, str], str]:
    hint, validation = _kvk(raw_id)
    raw = "" if raw_id is None else str(raw_id).strip()
    return ({
        "original_name": name, "country": "Nederland",
        "nl_evidence": "TenderNed gegunde onderneming" + (f"; plaats={city}" if city else ""),
        "employees_raw": "", "employees_date": "", "employees_scope": "",
        "website": "", "sector": "aanbestedingsleverancier",
        "source_kvk_hint": hint, "source_id": SOURCE_ID, "source_url": source_url,
        "fetched_at": fetched_at, "source_row": source_row,
        "evidence_reference": evidence_name, "source_registration_raw": raw,
        "registration_validation_status": validation, "source_legal_form": "",
        "source_status": "GUNNING_GEPUBLICEERD_NOT_KVK_VERIFIED", "candidate_layer": "raw",
    }, validation)


def _parse(
    run: Run, xlsx: Path, json_path: Path, latest_year: int, urls: tuple[str, str],
) -> tuple[Path, Path, dict[str, object]]:
    _safe_xlsx(xlsx)
    candidates = run.artifact_path("02", f"source_{SOURCE_ID}_companies", "csv")
    rejected = run.artifact_path("02", f"{SOURCE_ID}_rejected", "csv")
    candidate_temp = candidates.with_name(f".{candidates.name}.tmp")
    rejected_temp = rejected.with_name(f".{rejected.name}.tmp")
    counts: Counter[str] = Counter()
    unique_kvk: set[str] = set()
    xlsx_years: set[int] = set()
    fetched_at = datetime.now(UTC).isoformat()
    try:
        with (
            candidate_temp.open("w", encoding="utf-8", newline="") as candidate_handle,
            rejected_temp.open("w", encoding="utf-8", newline="") as rejected_handle,
        ):
            writer = csv.DictWriter(candidate_handle, fieldnames=RAW_HEADERS, delimiter="\t")
            review = csv.DictWriter(rejected_handle, fieldnames=REJECTED_HEADERS, delimiter="\t")
            writer.writeheader()
            review.writeheader()

            def add(name: str, raw_id: object, country: str, city: str,
                    source_row: str, source_url: str, evidence_name: str) -> None:
                counts["supplier_observations"] += 1
                raw = "" if raw_id is None else str(raw_id).strip()
                if not name or country.casefold() not in {"nederland", "nl"}:
                    counts["excluded_records"] += 1
                    reason = "MISSING_NAME" if not name else "NON_NL_OR_UNKNOWN_COUNTRY"
                    review.writerow({"source_row": source_row, "original_name": name,
                                     "source_registration_raw": raw, "country": country, "reason": reason})
                    counts[reason] += 1
                    return
                candidate, validation = _candidate(
                    name=name, raw_id=raw_id, city=city, source_row=source_row,
                    source_url=source_url, evidence_name=evidence_name, fetched_at=fetched_at,
                )
                writer.writerow(candidate)
                counts["candidate_records"] += 1
                counts[f"{validation.casefold()}_kvk_records"] += 1
                if validation == "VALID":
                    unique_kvk.add(candidate["source_kvk_hint"])
                else:
                    review.writerow({"source_row": source_row, "original_name": name,
                                     "source_registration_raw": raw, "country": country,
                                     "reason": f"{validation}_KVK"})
                    counts["review_records"] += 1

            workbook = load_workbook(xlsx, read_only=True, data_only=True)
            try:
                sheet = workbook["OpenData sheet"]
                rows = sheet.iter_rows(values_only=True)
                header = next(rows)
                index = {str(value): position for position, value in enumerate(header)}
                if not REQUIRED_XLSX.issubset(index):
                    raise HarvestError("TenderNed-XLSX mist verplichte kolommen")
                for row_number, row in enumerate(rows, 2):
                    counts["xlsx_publication_rows"] += 1
                    date = str(row[index["Publicatiedatum"]] or "")
                    if not re.match(r"^\d{4}-\d{2}-\d{2}$", date):
                        raise HarvestError("TenderNed-XLSX bevat ongeldige publicatiedatum")
                    xlsx_years.add(int(date[:4]))
                    if int(date[:4]) >= latest_year:
                        counts["xlsx_latest_year_superseded"] += 1
                        continue
                    name = str(row[index["Naam gegunde onderneming"]] or "").strip()
                    if not name:
                        counts["xlsx_without_awarded_name"] += 1
                        continue
                    url = str(row[index["URL TenderNed"]] or "")
                    if not url.startswith("https://www.tenderned.nl/"):
                        url = urls[0]
                    add(name, row[index["ON kvknummer"]], str(row[index["ON land"]] or "").strip(),
                        str(row[index["ON plaats"]] or "").strip(), f"xlsx:{row_number}",
                        url, xlsx.name)
            finally:
                workbook.close()
            if xlsx_years != set(range(2021, latest_year + 1)):
                raise HarvestError("TenderNed-XLSX mist een jaar in de archiefperiode 2021–laatste JSON-jaar")

            with json_path.open(encoding="utf-8") as handle:
                document = json.load(handle)
            releases = document.get("releases") if isinstance(document, dict) else None
            if not isinstance(releases, list) or not releases:
                raise HarvestError("TenderNed-JSON mist publicaties")
            for release_index, release in enumerate(releases):
                if not isinstance(release, dict):
                    raise HarvestError("TenderNed-JSON bevat ongeldige publicatie")
                release_date = str(release.get("date") or "")
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", release_date) or int(release_date[:4]) != latest_year:
                    raise HarvestError("TenderNed-JSON past niet bij laatste XLSX-jaar")
                counts["json_release_rows"] += 1
                awarded = {
                    (str(supplier.get("id") or ""), str(supplier.get("name") or ""))
                    for award in (release.get("awards") or [])
                    for supplier in (award.get("suppliers") or [])
                    if isinstance(supplier, dict)
                }
                release_id = str(release.get("id") or "")
                if not release_id.isascii() or not release_id.isdecimal():
                    raise HarvestError("TenderNed-JSON bevat ongeldige publicatie-ID")
                publication_url = f"https://www.tenderned.nl/tenderned-tap/aankondigingen/{release_id}"
                for party_index, party in enumerate(release.get("parties", [])):
                    if not isinstance(party, dict) or "supplier" not in party.get("roles", []):
                        continue
                    counts["json_supplier_party_rows"] += 1
                    key = (str(party.get("id") or ""), str(party.get("name") or ""))
                    if key not in awarded:
                        raise HarvestError("TenderNed-leverancier mist gunningsrelatie")
                    address = party.get("address") or {}
                    if not isinstance(address, dict):
                        raise HarvestError("TenderNed-leverancier heeft ongeldige adresstructuur")
                    add(str(party.get("name") or "").strip(), party.get("id"),
                        str(address.get("countryName") or "").strip(),
                        str(address.get("locality") or "").strip(),
                        f"json:{release_index}:{party_index}", publication_url, json_path.name)
            candidate_handle.flush()
            os.fsync(candidate_handle.fileno())
            rejected_handle.flush()
            os.fsync(rejected_handle.fileno())
        if counts["supplier_observations"] != counts["candidate_records"] + counts["excluded_records"]:
            raise HarvestError("TenderNed-innamepartitie sluit niet")
        os.replace(candidate_temp, candidates)
        os.replace(rejected_temp, rejected)
    finally:
        candidate_temp.unlink(missing_ok=True)
        rejected_temp.unlink(missing_ok=True)
    counts["unique_valid_kvk"] = len(unique_kvk)
    return candidates, rejected, {
        "source_id": SOURCE_ID, "scope": "FULL_ARCHIVE", "configured_limit": None,
        "status_semantics": "SOURCE_DATA_NOT_KVK_VERIFICATION", "latest_json_year": latest_year,
        "observed_xlsx_years": sorted(xlsx_years),
        "counts": dict(counts),
        "count_closure": {"supplier_partition": "CLOSED"},
    }


def _reusable(run: Run, config: object, kinds: dict[str, str]) -> bool:
    if not isinstance(config, dict):
        return False
    paths = config.get("output_paths")
    digests = config.get("output_sha256")
    if not isinstance(paths, dict) or not isinstance(digests, dict):
        return False
    for key, kind in kinds.items():
        path = run.latest_artifact("02", kind)
        if (path is None or not path.is_file() or str(path.relative_to(run.path)) != paths.get(key)
                or sha256(path) != digests.get(key)):
            return False
    return True


def collect_tenderned(
    run: Run, xlsx: Path | None = None, json_file: Path | None = None,
    refresh: bool = False,
) -> tuple[Path, Path, Path, Path]:
    """Download/hergebruik officiële snapshots; geen KVK-verzoek of impliciete bronverruiming."""
    with run.lock():
        if (xlsx is None) != (json_file is None):
            raise HarvestError("TenderNed vereist beide lokale bestanden of geen van beide")
        if not run.latest_artifact("01", "sources_inventory"):
            from company_harvest.sources import discover
            discover(run)
        kinds = {
            "source": f"source_{SOURCE_ID}", "rejected": f"{SOURCE_ID}_rejected",
            "report": f"{SOURCE_ID}_ingest_report", "md": f"{SOURCE_ID}_ingest_report_md",
            "xlsx": f"evidence_{SOURCE_ID}_xlsx", "json": f"evidence_{SOURCE_ID}_json",
        }
        config_key = f"{SOURCE_ID}_collect"
        config = run.metadata().get("runtime_config", {}).get(config_key)
        if xlsx is not None and json_file is not None:
            if not xlsx.expanduser().is_file() or not json_file.expanduser().is_file():
                raise HarvestError("TenderNed-lokale XLSX of JSON ontbreekt")
            local_hashes = {
                "xlsx": sha256(xlsx.expanduser().resolve()),
                "json": sha256(json_file.expanduser().resolve()),
            }
        else:
            local_hashes = None
        mode = "local-files" if local_hashes is not None else "download"
        if mode == "download":
            kinds["page"] = f"evidence_{SOURCE_ID}_page"
        if (not refresh and _reusable(run, config, kinds) and isinstance(config, dict)
                and config.get("mode") == mode and config.get("input_sha256") == local_hashes):
            run.log("INFO", "tenderned_collect_reused")
            return tuple(run.latest_artifact("02", kinds[key]) for key in (
                "source", "rejected", "report", "md"
            ))  # type: ignore[return-value]
        if shutil.disk_usage(run.path).free < 2 * MAX_FILE_BYTES + 512 * 1024 * 1024:
            raise HarvestError("onvoldoende vrije ruimte voor TenderNed-download en verwerking")
        evidence_dir = run.path / "evidence"
        excel_evidence = evidence_dir / f"{timestamp()}_02_tenderned.xlsx"
        json_evidence = evidence_dir / f"{timestamp()}_02_tenderned.json"
        page_evidence: Path | None = None
        acquisition: dict[str, object]
        if local_hashes is None:
            page_evidence = evidence_dir / f"{timestamp()}_02_tenderned_page.html"
            page_meta = _download(PAGE_URL, page_evidence, MAX_PAGE_BYTES)
            urls = _select_downloads(page_evidence.read_text(encoding="utf-8"))
            excel_meta = _download(urls[0], excel_evidence, MAX_FILE_BYTES)
            json_meta = _download(urls[1], json_evidence, MAX_FILE_BYTES)
            latest_year = urls[2]
            source_urls = (urls[0], urls[1])
            acquisition = {"mode": mode, "page": page_meta, "xlsx": excel_meta, "json": json_meta}
        else:
            assert xlsx is not None and json_file is not None
            _local_copy(xlsx, excel_evidence, MAX_FILE_BYTES)
            _local_copy(json_file, json_evidence, MAX_FILE_BYTES)
            with json_evidence.open(encoding="utf-8") as handle:
                json_data = json.load(handle)
            releases = json_data.get("releases") if isinstance(json_data, dict) else None
            if not isinstance(releases, list) or not releases:
                raise HarvestError("TenderNed-JSON mist publicaties")
            first_date = str(releases[0].get("date") or "") if isinstance(releases[0], dict) else ""
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", first_date):
                raise HarvestError("TenderNed-JSON mist geldige publicatiedatum")
            latest_year = int(first_date[:4])
            source_urls = (PAGE_URL, PAGE_URL)
            acquisition = {"mode": mode, "input_sha256": local_hashes}
        try:
            candidate_path, rejected_path, report = _parse(
                run, excel_evidence, json_evidence, latest_year, source_urls,
            )
        except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, json.JSONDecodeError) as exc:
            raise HarvestError(f"TenderNed-verwerking mislukt: {type(exc).__name__}") from exc
        run.invalidate_from(3, "tenderned_collection_changed")
        evidence_paths = [excel_evidence, json_evidence] + ([page_evidence] if page_evidence else [])
        report["evidence_artifacts"] = [
            {"path": str(path.relative_to(run.path)), "sha256": sha256(path), "size": path.stat().st_size}
            for path in evidence_paths
        ]
        report["candidate_artifact"] = {
            "path": str(candidate_path.relative_to(run.path)),
            "sha256": sha256(candidate_path), "size": candidate_path.stat().st_size,
        }
        report["acquisition"] = acquisition
        if mode == "local-files":
            report["scope"] = "LOCAL_UNVERIFIED"
        report["provenance_status"] = (
            "OFFICIAL_HOST_DOWNLOAD" if mode == "download"
            else "USER_SUPPLIED_OFFICIAL_FORMAT_UNVERIFIED"
        )
        report_path = run.artifact_path("02", f"{SOURCE_ID}_ingest_report", "json")
        atomic_write(report_path, json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        md_path = run.artifact_path("02", f"{SOURCE_ID}_ingest_report", "md")
        counts = cast(dict[str, int], report["counts"])
        atomic_write(md_path, "\n".join([
            "# TenderNed-inname", "", f"- Scope: 2021–{latest_year} (laatste jaar JSON)",
            f"- Kandidaten: {counts['candidate_records']}",
            f"- Unieke achtcijferige KVK-bronhints: {counts['unique_valid_kvk']}",
            f"- Buitenlandse/naamloze leveranciersregels: {counts['excluded_records']}",
            "- Status: bronhint, niet KVK-geverifieerd", "",
        ]))
        entries = [
            (excel_evidence, "02", kinds["xlsx"], "COMPLETE"),
            (json_evidence, "02", kinds["json"], "COMPLETE"),
            (candidate_path, "02", kinds["source"], "COMPLETE"),
            (rejected_path, "02", kinds["rejected"], "COMPLETE"),
            (report_path, "02", kinds["report"], "COMPLETE"),
            (md_path, "02", kinds["md"], "COMPLETE"),
        ]
        if page_evidence:
            entries.append((page_evidence, "02", kinds["page"], "COMPLETE"))
        run.register_artifact_set(entries)
        rows = read_catalog(run, migrate=True)
        for row in rows:
            if row["source_id"] == SOURCE_ID:
                row["status"] = "COLLECTED"
                row["live_measurement_status"] = (
                    "MEASURED_LIVE_EVIDENCE" if mode == "download"
                    else "MEASURED_LOCAL_EVIDENCE"
                )
                row["measured_count"] = str(counts["candidate_records"])
        inventory = run.artifact_path("01", "sources_inventory_tenderned_measured", "csv")
        from company_harvest.core import write_tsv
        write_tsv(inventory, SOURCE_HEADERS, rows)
        run.register_artifact(inventory, "01", "sources_inventory")
        paths = {
            "source": candidate_path, "rejected": rejected_path, "report": report_path,
            "md": md_path, "xlsx": excel_evidence, "json": json_evidence,
        }
        if page_evidence:
            paths["page"] = page_evidence
        run.record_config(config_key, {
            "mode": mode, "input_sha256": local_hashes,
            "output_paths": {key: str(path.relative_to(run.path)) for key, path in paths.items()},
            "output_sha256": {key: sha256(path) for key, path in paths.items()},
            "evidence_sha256": [sha256(path) for path in evidence_paths],
        })
        run.log("INFO", "tenderned_collect_completed", candidates=counts["candidate_records"],
                valid_kvk=counts["valid_kvk_records"])
        run.update_status("IN_PROGRESS", "02")
        return candidate_path, rejected_path, report_path, md_path
