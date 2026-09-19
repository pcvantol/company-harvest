"""Zelfstandige offline MERGE_LISTS-workflow."""

from __future__ import annotations

import csv
import io
import json
import shutil
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from company_lookup.core import HarvestError, Run, normalize_name, sha256, validate_kvk, write_tsv
from company_lookup.workflow import write_xlsx

ALIASES = {
    "name": ("Bedrijfsnaam", "naam"),
    "kvk": ("KVK-nummer", "kvkNummer", "kvk_nummer"),
    "sector": ("sector",),
    "city": ("plaats", "Vestigingsplaats"),
    "website": ("website",),
    "active": ("actief",),
}


@dataclass(frozen=True)
class InputOptions:
    delimiter: str | None = None
    encoding: str = "utf-8-sig"
    sheet: str | None = None
    name_column: str | None = None
    kvk_column: str | None = None


def _detect_delimiter(sample: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t,;")
        return dialect.delimiter
    except csv.Error as exc:
        raise HarvestError("CSV-delimiter is ambigu; geef een expliciete override") from exc


def _read_csv(path: Path, options: InputOptions) -> tuple[list[str], list[tuple[int, dict[str, Any]]], str]:
    text = path.read_text(encoding=options.encoding)
    delimiter = options.delimiter or _detect_delimiter(text[:8192])
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    if not reader.fieldnames:
        raise HarvestError(f"header ontbreekt in {path}")
    return list(reader.fieldnames), [(index, dict(row)) for index, row in enumerate(reader, 2)], f"CSV delimiter={delimiter!r}"


def _read_xlsx(path: Path, options: InputOptions) -> tuple[list[str], list[tuple[int, dict[str, Any]]], str]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if options.sheet:
        if options.sheet not in workbook.sheetnames:
            raise HarvestError(f"werkblad ontbreekt: {options.sheet}")
        sheet = workbook[options.sheet]
    elif len(workbook.sheetnames) == 1:
        sheet = workbook[workbook.sheetnames[0]]
    elif "Bedrijven" in workbook.sheetnames:
        sheet = workbook["Bedrijven"]
    else:
        raise HarvestError("XLSX-werkblad is ambigu; geef een expliciete override")
    iterator = sheet.iter_rows(values_only=False)
    try:
        first = next(iterator)
    except StopIteration as exc:
        raise HarvestError(f"lege XLSX: {path}") from exc
    headers = [str(cell.value or "").strip() for cell in first]
    if not all(headers) or len(headers) != len(set(headers)):
        raise HarvestError("XLSX-header bevat lege of dubbele kolommen")
    rows: list[tuple[int, dict[str, Any]]] = []
    for index, cells in enumerate(iterator, 2):
        row: dict[str, Any] = {}
        for header, cell in zip(headers, cells, strict=False):
            value = cell.value
            if cell.data_type == "f":
                value = None
            row[header] = value
        rows.append((index, row))
    return headers, rows, f"XLSX sheet={sheet.title}"


def _mapping(headers: list[str], options: InputOptions) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {}
    for canonical, aliases in ALIASES.items():
        explicit = options.name_column if canonical == "name" else options.kvk_column if canonical == "kvk" else None
        matches = [header for header in headers if header in aliases]
        if explicit:
            if explicit not in headers:
                raise HarvestError(f"expliciete kolom ontbreekt: {explicit}")
            mapping[canonical] = explicit
        elif len(matches) == 1:
            mapping[canonical] = matches[0]
        elif len(matches) > 1:
            raise HarvestError(f"ambigue kolommen voor {canonical}: {matches}")
        else:
            mapping[canonical] = None
    if not mapping["name"] or not mapping["kvk"]:
        raise HarvestError("verplichte naam- of KVK-kolom ontbreekt")
    return mapping


def _load(path: Path, side: str, options: InputOptions, snapshot: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    if not path.is_file():
        raise HarvestError(f"inputbestand ontbreekt: {path}")
    shutil.copyfile(path, snapshot)
    headers, raw_rows, format_detail = _read_xlsx(path, options) if path.suffix.lower() == ".xlsx" else _read_csv(path, options)
    mapping = _mapping(headers, options)
    valid, rejected = [], []
    for row_number, row in raw_rows:
        raw_name = row.get(str(mapping["name"]))
        name = str(raw_name or "").strip()
        try:
            number = validate_kvk(row.get(str(mapping["kvk"])))
            if not name:
                raise ValueError("bedrijfsnaam ontbreekt")
        except ValueError as exc:
            rejected.append({"input_side": side, "file": path.name, "row": str(row_number), "original_name": name, "original_kvk": str(row.get(str(mapping["kvk"])) or ""), "reason": str(exc)})
            continue
        known = {key: (row.get(column) if column else "") for key, column in mapping.items()}
        extras = {header: row.get(header) for header in headers if header not in set(value for value in mapping.values() if value)}
        valid.append({"side": side, "file": path.name, "file_sha256": sha256(path), "row": row_number, "name": name, "normalized_name": normalize_name(name), "kvk": number, **known, "extras_json": json.dumps(extras, ensure_ascii=False, default=str, separators=(",", ":"))})
    return valid, rejected, {"path": str(path.resolve()), "sha256": sha256(path), "format": format_detail, "headers": headers, "mapping": mapping, "read": len(raw_rows), "valid": len(valid), "rejected": len(rejected)}


def merge_lists(run: Run, left: Path, right: Path, policy: str, left_options: InputOptions = InputOptions(), right_options: InputOptions = InputOptions()) -> list[Path]:
    if policy not in {"exclude", "prefer-left", "prefer-right"}:
        raise HarvestError("ongeldig conflictbeleid")
    with run.lock():
        from company_lookup.core import timestamp

        left_snapshot = run.path / "snapshots" / f"{timestamp()}_09_left_{left.name}"
        right_snapshot = run.path / "snapshots" / f"{timestamp()}_09_right_{right.name}"
        left_rows, left_rejected, left_meta = _load(left, "left", left_options, left_snapshot)
        right_rows, right_rejected, right_meta = _load(right, "right", right_options, right_snapshot)
        run.register_artifact(left_snapshot, "09", "input_left_snapshot")
        run.register_artifact(right_snapshot, "09", "input_right_snapshot")
        groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in left_rows + right_rows:
            groups[row["kvk"]].append(row)
        merged: list[dict[str, Any]] = []
        conflicts: list[dict[str, Any]] = []
        field_conflicts: list[dict[str, Any]] = []
        for kvk, rows in groups.items():
            names = {row["normalized_name"] for row in rows}
            selected_name: str | None = None
            reason = ""
            if len(names) == 1:
                selected_name = next(row["name"] for row in rows if row["side"] == "left") if any(row["side"] == "left" for row in rows) else rows[0]["name"]
            elif policy != "exclude":
                side = policy.removeprefix("prefer-")
                side_names = {row["normalized_name"] for row in rows if row["side"] == side}
                if len(side_names) == 1:
                    selected_name = next(row["name"] for row in rows if row["side"] == side)
                    reason = "RESOLVED_BY_EXPLICIT_PREFERENCE"
                else:
                    reason = "UNRESOLVED_WITHIN_PREFERRED_SOURCE"
            else:
                reason = "NAME_CONFLICT_EXCLUDED"
            if len(names) > 1:
                group_id = f"NC-{kvk}"
                for row in rows:
                    conflicts.append({"conflict_group_id": group_id, "KVK-nummer": kvk, "original_name": row["name"], "input_side": row["side"], "file": row["file"], "file_sha256": row["file_sha256"], "sheet_row": row["row"], "reason": reason, "conflict_policy": policy, "chosen_name": selected_name or "", "included": str(bool(selected_name)).lower()})
            if not selected_name:
                continue
            output: dict[str, Any] = {"Bedrijfsnaam": selected_name, "KVK-nummer": kvk}
            for field, label in (("sector", "Sector"), ("city", "Vestigingsplaats"), ("website", "Website"), ("active", "Actieve status")):
                values = {str(row.get(field) or "").strip() for row in rows if str(row.get(field) or "").strip()}
                if len(values) == 1:
                    output[label] = next(iter(values))
                elif len(values) > 1:
                    output[label] = "UNKNOWN"
                    for row in rows:
                        if str(row.get(field) or "").strip():
                            field_conflicts.append({"KVK-nummer": kvk, "field": field, "value": row[field], "input_side": row["side"], "file": row["file"], "row": row["row"]})
                else:
                    output[label] = ""
            output["source_rows_json"] = json.dumps([{"side": row["side"], "file": row["file"], "sha256": row["file_sha256"], "row": row["row"], "name": row["name"], "extras": json.loads(row["extras_json"])} for row in rows], ensure_ascii=False, separators=(",", ":"))
            output["verification_status"] = "UNCONFIRMED_IMPORTED"
            merged.append(output)
        merged.sort(key=lambda row: (normalize_name(row["Bedrijfsnaam"]), row["KVK-nummer"]))
        rejected = left_rejected + right_rejected
        paths = [run.artifact_path("09", "companies_merged", "csv"), run.artifact_path("09", "companies_merged", "xlsx"), run.artifact_path("09", "companies_merged_full", "csv"), run.artifact_path("09", "conflicts", "csv"), run.artifact_path("09", "field_conflicts", "csv"), run.artifact_path("09", "rejected_rows", "csv"), run.artifact_path("09", "merge_report", "md")]
        minimal = ["Bedrijfsnaam", "KVK-nummer"]
        full_headers = minimal + ["Sector", "Vestigingsplaats", "Website", "Actieve status", "source_rows_json", "verification_status"]
        write_tsv(paths[0], minimal, merged)
        write_xlsx(paths[1], minimal, merged)
        write_tsv(paths[2], full_headers, merged)
        write_tsv(paths[3], ["conflict_group_id", "KVK-nummer", "original_name", "input_side", "file", "file_sha256", "sheet_row", "reason", "conflict_policy", "chosen_name", "included"], conflicts)
        write_tsv(paths[4], ["KVK-nummer", "field", "value", "input_side", "file", "row"], field_conflicts)
        write_tsv(paths[5], ["input_side", "file", "row", "original_name", "original_kvk", "reason"], rejected)
        report = _report(left_meta, right_meta, policy, merged, groups, rejected, conflicts, field_conflicts)
        paths[6].write_text(report, encoding="utf-8")
        for path, kind in zip(paths, ("merged_csv", "merged_xlsx", "merged_full", "conflicts", "field_conflicts", "rejected_rows", "merge_report"), strict=True):
            run.register_artifact(path, "09", kind)
        run.update_status("COMPLETE_WITH_WARNINGS" if conflicts or rejected else "COMPLETE", "09")
        return paths


def _report(left: dict[str, Any], right: dict[str, Any], policy: str, merged: list[dict[str, Any]], groups: dict[str, list[dict[str, Any]]], rejected: list[dict[str, Any]], conflicts: list[dict[str, Any]], field_conflicts: list[dict[str, Any]]) -> str:
    valid = left["valid"] + right["valid"]
    read = left["read"] + right["read"]
    if valid + len(rejected) != read:
        raise HarvestError("interne telling merge-invoer klopt niet", 7)
    conflict_groups = len({row["conflict_group_id"] for row in conflicts})
    return "\n".join([
        "# Merge-rapport", "", f"- Conflictbeleid: `{policy}`", f"- Links: `{json.dumps(left, ensure_ascii=False)}`", f"- Rechts: `{json.dumps(right, ensure_ascii=False)}`", f"- Gelezen regels: {read}", f"- Geldige regels: {valid}", f"- Afgewezen regels: {len(rejected)}", f"- Unieke geldige KVK-nummers: {len(groups)}", f"- Overtollige duplicaatregels: {valid - len(groups)}", f"- Naamconflictgroepen: {conflict_groups}", f"- Naamconflictregels: {len(conflicts)}", f"- Veldconflicten: {len(field_conflicts)}", f"- Definitieve ondernemingen: {len(merged)}", "", "Geïmporteerde nummers zijn syntactisch gevalideerd maar niet opnieuw bij KVK gecontroleerd. De uitvoer kan inactieve of onbevestigde ondernemingen bevatten.", "",
    ])
